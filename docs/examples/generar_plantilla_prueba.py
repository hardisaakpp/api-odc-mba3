"""Genera la plantilla Excel de prueba (formato v1.0: hojas 'cabecera' y 'detalle') para
"Cargar Plantilla" del módulo ODC SAP. Datos ficticios, solo para ambiente local."""
import sys
from datetime import datetime

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

SALIDA = sys.argv[1] if len(sys.argv) > 1 else "plantilla_pedido_prueba.xlsx"
FUENTE = "Arial"

F_BASE = Font(name=FUENTE, size=10)
F_HEADER = Font(name=FUENTE, size=10, bold=True, color="FFFFFF")
F_CODIGOS = Font(name=FUENTE, size=9, italic=True, color="404040")
F_INPUT = Font(name=FUENTE, size=10, color="0000FF")
F_TITULO = Font(name=FUENTE, size=14, bold=True)
F_SUBT = Font(name=FUENTE, size=11, bold=True)

FILL_HEADER = PatternFill("solid", fgColor="1F4E78")
FILL_BODEGA = PatternFill("solid", fgColor="7F6000")
FILL_TIENDA = PatternFill("solid", fgColor="375623")
FILL_CALC = PatternFill("solid", fgColor="595959")
FILL_CODIGOS = PatternFill("solid", fgColor="E7E6E6")
FILL_INPUT = PatternFill("solid", fgColor="FFF2CC")
BORDE = Border(*(Side(style="thin", color="BFBFBF"),) * 4)

# ---------------------------------------------------------------- datos de prueba
CABECERA = [
    ("CODIGO PEDIDO", "PRUEBA-0001", "Identificador único del pedido. CAMBIARLO en cada carga (el sistema rechaza duplicados)."),
    ("FECHA DE PEDIDO", datetime(2026, 9, 24), "Fecha del pedido."),
    ("FECHA DE ENTREGA", "INMEDIATO", "Tipo o fecha de entrega."),
    ("FORMA DE PAGO", "120 DIAS", "Condición de pago."),
    ("MENSAJE", "Pedido de prueba - ambiente local", "Texto libre (máx. 500 caracteres)."),
    ("NOMBRE PROVEEDOR", "TEXTILES EL RAYO S.A.", "Proveedor del pedido."),
    ("COMPRADOR NOMBRE", "PRUEBAS SCNEXUS", "Nombre del comprador."),
    ("COD COMPRADOR", "801", "Código del comprador."),
    ("ORIGEN TRASLADO", "RL-CD", "Bodega origen de los traslados (RL-CD, OUT-CD o RG-CD)."),
    ("SERIE", "10131", "Serie del documento."),
]

BODEGAS = [("BODEGA ROLAND", "RL-CD"), ("BODEGA DESCUENTON", "OUT-CD"), ("BODEGA REGIONS", "RG-CD")]
TIENDAS = [
    ("ROLAND PSC", "TRL-PSC"), ("ROLAND PSP", "TRL-PSP"), ("ROLAND PCH", "TRL-PCH"),
    ("DESCUENTON LLG", "TOUT-LLG"), ("DESCUENTON QNT", "TOUT-QNT"), ("REGIONS MLS", "RG-MLS"),
]

# codigo, descripcion, pack, costo sin IVA (unitario; el backend advierte si costo*1.15 > PVP), PVP, [bodegas RL/OUT/RG], [tiendas en orden de TIENDAS]
PRODUCTOS = [
    ("R9001.101.NEG", "CAMISETA BASICA PRUEBA NEGRO", "X6", 5.60, 12.99, [2, 1, None], [3, 2, 2, 1, 1, 2]),
    ("R9001.101.BLA", "CAMISETA BASICA PRUEBA BLANCO", "X6", 5.60, 12.99, [2, None, None], [2, 2, 1, 1, None, 1]),
    ("R9002.205.AZM", "JEAN SLIM PRUEBA AZUL MEDIO", "X3", 17.50, 39.99, [1, None, 1], [2, 1, 1, None, None, 1]),
    ("R9003.310.GRS", "BUZO CAPUCHA PRUEBA GRIS", "X2", 15.20, 34.99, [None, None, None], [1, 1, 0.5, None, None, None]),
    ("R9004.050.SUR", "MEDIAS PACK PRUEBA SURTIDO", "X12", 2.10, 4.99, [3, 2, 1], [4, 4, 3, 2, 2, 2]),
    ("R9005.412.ROJ", "CHOMPA TEJIDA PRUEBA ROJO", "X4", 13.00, 29.99, [1, None, None], [1, 1, 1, 1, 1, None]),
    ("R9006.600.BEI", "PANTALON DRILL PRUEBA BEIGE", "X3", 14.30, 32.99, [None, 1, None], [2, None, 1, 1, 2, 1]),
    ("R9007.777.VRD", "GORRA PRUEBA VERDE", "X12", 4.20, 9.99, [2, None, None], [None] * 6),
]
PACKS_VALIDOS = ["X1", "X2", "X3", "X4", "X5", "X6", "X7", "X8", "X9", "X12"]  # tabla pack_conversion


def ean13(base12: str) -> str:
    """Código EAN-13 con dígito verificador (prefijo 786 = Ecuador; ficticio)."""
    suma = sum(int(d) * (3 if i % 2 else 1) for i, d in enumerate(base12))
    return base12 + str((10 - suma % 10) % 10)


def estilo(celda, font=F_BASE, fill=None, fmt=None, align=None):
    celda.font = font
    celda.border = BORDE
    if fill:
        celda.fill = fill
    if fmt:
        celda.number_format = fmt
    if align:
        celda.alignment = align


wb = Workbook()

# ---------------------------------------------------------------- hoja cabecera
ws = wb.active
ws.title = "cabecera"
for col, titulo in enumerate(("CAMPO", "VALOR", "DESCRIPCION"), start=1):
    estilo(ws.cell(1, col, titulo), F_HEADER, FILL_HEADER)
for fila, (campo, valor, desc) in enumerate(CABECERA, start=2):
    estilo(ws.cell(fila, 1, campo), Font(name=FUENTE, size=10, bold=True))
    fmt = "yyyy-mm-dd" if isinstance(valor, datetime) else "@"
    estilo(ws.cell(fila, 2, valor), F_INPUT, FILL_INPUT, fmt, Alignment(horizontal="left"))
    estilo(ws.cell(fila, 3, desc), Font(name=FUENTE, size=9, italic=True, color="595959"))
ws["B2"].comment = Comment("Debe ser único: el sistema rechaza un CODIGO PEDIDO que ya exista "
                           "(NUMERO_PEDIDO_DUPLICADO). Usa PRUEBA-0002, PRUEBA-0003, ...", "SCNexus pruebas")
ws.column_dimensions["A"].width = 22
ws.column_dimensions["B"].width = 36
ws.column_dimensions["C"].width = 80
ws.freeze_panes = "A2"

# ---------------------------------------------------------------- hoja detalle
wd = wb.create_sheet("detalle")
cols_producto = ["Codigo Producto", "Codigo de Barras", "Descripcion SAP", "Pack",
                 "Costo sin IVA", "Costo Unitario", "Descuento", "PVP"]
c_bod0 = len(cols_producto) + 1                 # primera columna de bodegas (I)
c_tie0 = c_bod0 + len(BODEGAS)                  # primera columna de tiendas (L)
c_tie1 = c_tie0 + len(TIENDAS) - 1              # última columna de tiendas (Q)
c_total, c_subtotal = c_tie1 + 1, c_tie1 + 2    # R, S
L = get_column_letter

centrado = Alignment(horizontal="center", vertical="center", wrap_text=True)
for i, nombre in enumerate(cols_producto, start=1):
    estilo(wd.cell(1, i, nombre), F_HEADER, FILL_HEADER, align=centrado)
    estilo(wd.cell(2, i), F_CODIGOS, FILL_CODIGOS)
for i, (nombre, codigo) in enumerate(BODEGAS):
    estilo(wd.cell(1, c_bod0 + i, nombre), F_HEADER, FILL_BODEGA, align=centrado)
    estilo(wd.cell(2, c_bod0 + i, codigo), F_CODIGOS, FILL_CODIGOS, align=centrado)
for i, (nombre, codigo) in enumerate(TIENDAS):
    estilo(wd.cell(1, c_tie0 + i, nombre), F_HEADER, FILL_TIENDA, align=centrado)
    estilo(wd.cell(2, c_tie0 + i, codigo), F_CODIGOS, FILL_CODIGOS, align=centrado)
for col, nombre in ((c_total, "TOTAL"), (c_subtotal, "SUBTOTAL")):
    estilo(wd.cell(1, col, nombre), F_HEADER, FILL_CALC, align=centrado)
    estilo(wd.cell(2, col), F_CODIGOS, FILL_CODIGOS)

wd["A1"].comment = Comment("Fila 1 = nombres de columna. Fila 2 = códigos de bodega/tienda "
                           "(el sistema identifica cada tienda por el código de la fila 2). "
                           "Productos desde la fila 3.", "SCNexus pruebas")
wd["D1"].comment = Comment("Pack SAP del producto (X1, X2, X3, X4, X5, X6, X7, X8, X9, X12). "
                           "Debe coincidir con el pack registrado en SAP.", "SCNexus pruebas")
wd["F1"].comment = Comment("Calculado: Costo sin IVA / (12 / n), con n = número del pack "
                           "(misma regla que la tabla pack_conversion del backend).", "SCNexus pruebas")
wd["H1"].comment = Comment("Debe coincidir con el Precio Minorista registrado en SAP.", "SCNexus pruebas")
wd[f"{L(c_total)}1"].comment = Comment("Bodegas + tiendas. El backend valida que coincida con la suma "
                                       "de cantidades de la fila.", "SCNexus pruebas")
wd[f"{L(c_subtotal)}1"].comment = Comment("Solo tiendas. El backend valida que coincida con la suma de "
                                          "cantidades de tiendas de la fila.", "SCNexus pruebas")

for n, (codigo, desc, pack, costo, pvp, bodegas, tiendas) in enumerate(PRODUCTOS):
    r = 3 + n
    estilo(wd.cell(r, 1, codigo), F_INPUT, FILL_INPUT)
    estilo(wd.cell(r, 2, ean13(f"786900{n + 1:06d}")), F_INPUT, FILL_INPUT, "@")
    estilo(wd.cell(r, 3, desc), F_INPUT, FILL_INPUT)
    estilo(wd.cell(r, 4, pack), F_INPUT, FILL_INPUT, align=Alignment(horizontal="center"))
    estilo(wd.cell(r, 5, costo), F_INPUT, FILL_INPUT, "#,##0.00")
    estilo(wd.cell(r, 6, f'=IFERROR(E{r}/(12/VALUE(MID(D{r},2,3))),"")'), fmt="#,##0.0000")
    estilo(wd.cell(r, 7, 0), F_INPUT, FILL_INPUT, "#,##0.00")
    estilo(wd.cell(r, 8, pvp), F_INPUT, FILL_INPUT, "#,##0.00")
    for i, cant in enumerate(bodegas + tiendas):
        estilo(wd.cell(r, c_bod0 + i, cant), F_INPUT, FILL_INPUT, "General",
               Alignment(horizontal="center"))
    estilo(wd.cell(r, c_total, f"=SUM({L(c_bod0)}{r}:{L(c_tie1)}{r})"),
           Font(name=FUENTE, size=10, bold=True), fmt="General", align=Alignment(horizontal="center"))
    estilo(wd.cell(r, c_subtotal, f"=SUM({L(c_tie0)}{r}:{L(c_tie1)}{r})"),
           Font(name=FUENTE, size=10, bold=True), fmt="General", align=Alignment(horizontal="center"))

ultima = 2 + len(PRODUCTOS)
dv = DataValidation(type="list", formula1='"' + ",".join(PACKS_VALIDOS) + '"', allow_blank=False,
                    showErrorMessage=True, errorTitle="Pack inválido",
                    error="Usa un pack de la lista (tabla pack_conversion).")
wd.add_data_validation(dv)
dv.add(f"D3:D{ultima + 50}")

anchos = {1: 17, 2: 16, 3: 34, 4: 7, 5: 11, 6: 11, 7: 10, 8: 9}
for col, ancho in anchos.items():
    wd.column_dimensions[L(col)].width = ancho
for col in range(c_bod0, c_subtotal + 1):
    wd.column_dimensions[L(col)].width = 12
wd.row_dimensions[1].height = 32
wd.freeze_panes = "D3"

# ---------------------------------------------------------------- hoja LEEME
wl = wb.create_sheet("LEEME")
lineas = [
    ("Plantilla de pedido de PRUEBA — módulo ODC SAP (SCNexus)", F_TITULO),
    ("Datos ficticios para el ambiente local. No usar en producción.", Font(name=FUENTE, size=10, italic=True)),
    ("", None),
    ("Qué editar", F_SUBT),
    ("• Celdas con fondo amarillo y texto azul: datos de entrada (cabecera, productos y cantidades).", F_BASE),
    ("• Celdas sin relleno: fórmulas (Costo Unitario, TOTAL, SUBTOTAL). No sobrescribirlas.", F_BASE),
    ("• cabecera!B2 (CODIGO PEDIDO): cambiarlo en cada carga; el sistema rechaza pedidos duplicados.", F_BASE),
    ("", None),
    ("Reglas del formato (validadas por el backend al cargar)", F_SUBT),
    ("• Las hojas deben llamarse exactamente 'cabecera' y 'detalle' (minúsculas).", F_BASE),
    ("• detalle: fila 1 = nombres de columna, fila 2 = códigos de bodega/tienda, productos desde la fila 3.", F_BASE),
    ("• Bodegas: códigos RL-CD (Roland), OUT-CD (Descuentón), RG-CD (Regions).", F_BASE),
    ("• Tiendas: el código de la fila 2 debe empezar con TRL-, TOUT- o RG- (ej. TRL-PSC).", F_BASE),
    ("• Para agregar una tienda: insertar una columna antes de TOTAL, con nombre en fila 1 y código en fila 2.", F_BASE),
    ("• TOTAL = bodegas + tiendas; SUBTOTAL = solo tiendas. Si no cuadran con las cantidades, la carga falla.", F_BASE),
    ("• Cantidades > 0; se permiten decimales (ej. 0.5). Celda vacía = sin distribución.", F_BASE),
    ("• Pack: X1, X2, X3, X4, X5, X6, X7, X8, X9 o X12. Pack y PVP deben coincidir con SAP.", F_BASE),
    ("", None),
    ("Casos que cubre esta plantilla", F_SUBT),
    ("• 8 productos, 3 bodegas y 6 tiendas de las 3 marcas.", F_BASE),
    ("• Producto solo con tiendas (R9003.310.GRS, con cantidad decimal 0.5).", F_BASE),
    ("• Producto solo con bodega (R9007.777.VRD).", F_BASE),
    ("• Distribuciones con celdas vacías y 5 packs distintos (X2, X3, X4, X6, X12).", F_BASE),
    ("", None),
    ("Importante", F_SUBT),
    ("• Al cargar, el sistema valida los códigos de producto, pack, PVP, tiendas y bodegas contra SAP.", F_BASE),
    ("  Sin SAP (o un simulador de SAP con estos mismos productos) la carga se detiene en esa validación.", F_BASE),
]
for i, (texto, fuente) in enumerate(lineas, start=1):
    c = wl.cell(i, 1, texto)
    c.font = fuente or F_BASE
wl.column_dimensions["A"].width = 110

wb.save(SALIDA)
print(f"OK {SALIDA}")

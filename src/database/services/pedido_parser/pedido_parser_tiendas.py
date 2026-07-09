"""
Parser especializado para tiendas y distribuciones de pedidos.

CHANGELOG:
- 2026-01-23: Extraído de pedido_parser.py
"""

import logging
import pandas as pd
import io
import re
import unicodedata
from typing import Dict, Any, List, Optional

from .pedido_parser_utils import PedidoParserUtils, normalize_column_name

logger = logging.getLogger(__name__)


class PedidoParserTiendas:
    """Parser para tiendas y distribuciones de pedidos."""

    def __init__(self):
        """Inicializa el parser con utilidades compartidas."""
        self.utils = PedidoParserUtils()

    @staticmethod
    def _buscar_hoja_movimiento(excel_file: pd.ExcelFile) -> Optional[str]:
        for sheet_name in excel_file.sheet_names:
            if str(sheet_name).upper().strip() == "MOVIMIENTO":
                return str(sheet_name)
        return None

    @staticmethod
    def _es_columna_codigo_producto(col_upper: str) -> bool:
        return ("CODIGO" in col_upper and "PRODUCTO" in col_upper) or col_upper in ["PRODUCTO", "SKU", "CODIGO"]

    @staticmethod
    def _es_columna_cantidad(col_upper: str) -> bool:
        return "CANTIDAD" in col_upper or col_upper in ["CANT", "QTY"]

    @staticmethod
    def _detectar_columnas_movimiento(df: pd.DataFrame) -> tuple:
        col_codigo = col_tienda = col_cantidad = None
        for col in df.columns:
            col_upper = str(col).upper().strip()
            if col_codigo is None and PedidoParserTiendas._es_columna_codigo_producto(col_upper):
                col_codigo = col
            if col_tienda is None and "TIENDA" in col_upper:
                col_tienda = col
            if col_cantidad is None and PedidoParserTiendas._es_columna_cantidad(col_upper):
                col_cantidad = col
        return col_codigo, col_tienda, col_cantidad

    @staticmethod
    def _parsear_cantidad_movimiento(
        cantidad_raw: Any, fila_num: int, resultado: Dict[str, Any]
    ) -> Optional[float]:
        if cantidad_raw is None or str(cantidad_raw).strip() == "":
            resultado["errores"].append(f"Fila {fila_num}: 'CANTIDAD' es obligatoria")
            return None
        try:
            cantidad = float(str(cantidad_raw).replace(",", "."))
            if cantidad < 0:
                resultado["errores"].append(f"Fila {fila_num}: 'CANTIDAD' no puede ser negativa")
                return None
            return cantidad
        except (ValueError, TypeError):
            resultado["errores"].append(f"Fila {fila_num}: 'CANTIDAD' no es un número válido")
            return None

    def _validar_tienda_catalogo(
        self,
        tienda_codigo: str,
        tienda_str: str,
        catalogo_tiendas: Optional[Dict[str, Dict[str, Any]]]
    ) -> tuple:
        tienda_nombre = tienda_str
        validacion_tienda: Dict[str, Any] = {"valida": True, "errores": [], "warnings": []}
        if not catalogo_tiendas:
            validacion_tienda["warnings"].append("No se encontró hoja 'Tiendas' para validar tiendas")
            return tienda_codigo, tienda_nombre, validacion_tienda
        for cod_cat, info_cat in catalogo_tiendas.items():
            nombre_cat_upper = str(info_cat.get("nombre", "")).upper().replace(" ", "")
            if tienda_codigo == cod_cat or tienda_codigo == nombre_cat_upper:
                tienda_codigo = cod_cat
                tienda_nombre = info_cat.get("nombre", tienda_str)
                if not info_cat.get("activa", False):
                    validacion_tienda["valida"] = False
                    validacion_tienda["errores"].append(
                        f"Tienda '{tienda_codigo}' ({tienda_nombre}) está INACTIVA"
                    )
                return tienda_codigo, tienda_nombre, validacion_tienda
        validacion_tienda["valida"] = False
        validacion_tienda["errores"].append(f"Tienda '{tienda_str}' no existe en el catálogo de tiendas")
        return tienda_codigo, tienda_nombre, validacion_tienda

    def _procesar_fila_movimiento(
        self,
        row: pd.Series,
        fila_num: int,
        col_codigo: str,
        col_tienda: str,
        col_cantidad: str,
        catalogo_tiendas: Optional[Dict[str, Dict[str, Any]]],
        productos_vistos: set,
        resultado: Dict[str, Any]
    ) -> None:
        codigo_producto = str(row.get(col_codigo, "")).strip()
        tienda_str = str(row.get(col_tienda, "")).strip()
        cantidad_raw = row.get(col_cantidad)

        if not codigo_producto:
            resultado["errores"].append(f"Fila {fila_num}: 'CODIGO PRODUCTO' es obligatorio")
            return
        if not tienda_str:
            resultado["errores"].append(f"Fila {fila_num}: 'TIENDAS' es obligatorio")
            return

        cantidad = self._parsear_cantidad_movimiento(cantidad_raw, fila_num, resultado)
        if cantidad is None:
            return

        codigo_producto = codigo_producto.upper()
        tienda_codigo = tienda_str.upper().replace(" ", "")
        tienda_codigo, tienda_nombre, validacion_tienda = self._validar_tienda_catalogo(
            tienda_codigo, tienda_str, catalogo_tiendas
        )

        resultado["datos"].append({
            "codigo_producto": codigo_producto,
            "tienda_codigo": tienda_codigo,
            "tienda_nombre": tienda_nombre,
            "cantidad": cantidad,
            "fila_origen": fila_num,
            "validacion": validacion_tienda
        })
        resultado["resumen"]["filas_validas"] += 1
        resultado["resumen"]["distribuciones_totales"] += 1

        if codigo_producto not in productos_vistos:
            productos_vistos.add(codigo_producto)
            resultado["resumen"]["productos_unicos"] += 1

    def leer_hoja_movimiento(
        self,
        file_content: bytes,
        catalogo_tiendas: Optional[Dict[str, Dict[str, Any]]] = None
    ) -> Dict[str, Any]:
        """
        Lee la hoja 'Movimiento' del Excel y retorna distribuciones por producto y tienda.

        Args:
            file_content: Contenido del archivo Excel en bytes
            catalogo_tiendas: Diccionario opcional con tiendas válidas

        Returns:
            Dict con distribuciones y resumen
        """
        resultado = {
            "resumen": {
                "filas_leidas": 0,
                "filas_validas": 0,
                "productos_unicos": 0,
                "distribuciones_totales": 0
            },
            "datos": [],
            "errores": []
        }

        try:
            excel_file = pd.ExcelFile(io.BytesIO(file_content))

            hoja_movimiento = self._buscar_hoja_movimiento(excel_file)
            if not hoja_movimiento:
                logger.warning("⚠️ No se encontró hoja 'Movimiento' en el Excel")
                return resultado

            df_movimiento = pd.read_excel(
                excel_file, sheet_name=hoja_movimiento, header=0, keep_default_na=False
            )
            df_movimiento.columns = [str(col).strip() for col in df_movimiento.columns]
            logger.info(f"📋 Hoja 'Movimiento' encontrada con {len(df_movimiento)} filas")

            col_codigo, col_tienda, col_cantidad = self._detectar_columnas_movimiento(df_movimiento)

            if not col_codigo:
                resultado["errores"].append("❌ No se encontró columna 'CODIGO PRODUCTO' en hoja Movimiento")
                return resultado
            if not col_tienda:
                resultado["errores"].append("❌ No se encontró columna 'TIENDAS' en hoja Movimiento")
                return resultado
            if not col_cantidad:
                resultado["errores"].append("❌ No se encontró columna 'CANTIDAD' en hoja Movimiento")
                return resultado

            productos_vistos: set = set()
            for idx, row in df_movimiento.iterrows():
                fila_num = (int(idx) if isinstance(idx, (int, float)) else 0) + 2
                resultado["resumen"]["filas_leidas"] += 1
                self._procesar_fila_movimiento(
                    row, fila_num, col_codigo, col_tienda, col_cantidad,
                    catalogo_tiendas, productos_vistos, resultado
                )

            logger.info(
                f"✅ Hoja 'Movimiento' procesada: {resultado['resumen']['filas_validas']} filas válidas, "
                f"{resultado['resumen']['productos_unicos']} productos únicos"
            )

        except Exception as e:
            logger.error(f"❌ Error al leer hoja 'Movimiento': {e}", exc_info=True)
            resultado["errores"].append(f"Error al leer hoja 'Movimiento': {str(e)}")

        return resultado

    def _validar_tienda_codigo(
        self,
        tienda_codigo: str,
        catalogo_tiendas: Optional[Dict[str, Dict[str, Any]]]
    ) -> tuple:
        validacion: Dict[str, Any] = {"valida": True, "errores": [], "warnings": []}
        if not catalogo_tiendas:
            validacion["warnings"].append("No se encontró hoja 'Tiendas' para validar el código")
            return tienda_codigo, validacion
        if tienda_codigo not in catalogo_tiendas:
            validacion["valida"] = False
            validacion["errores"].append(
                f"Código de tienda '{tienda_codigo}' no existe en hoja 'Tiendas'"
            )
            return tienda_codigo, validacion
        tienda_info = catalogo_tiendas[tienda_codigo]
        tienda_nombre = tienda_info.get("nombre", tienda_codigo)
        if not tienda_info.get("activa", False):
            validacion["valida"] = False
            validacion["errores"].append(
                f"Tienda '{tienda_codigo}' ({tienda_nombre}) está INACTIVA"
            )
        return tienda_nombre, validacion

    def _construir_distribucion_columna(
        self,
        col_name: str,
        col_upper: str,
        es_codigo_tienda: bool,
        cantidad_val: float,
        catalogo_tiendas: Optional[Dict[str, Dict[str, Any]]]
    ) -> Dict[str, Any]:
        if es_codigo_tienda:
            tienda_nombre, validacion = self._validar_tienda_codigo(col_upper, catalogo_tiendas)
            return {
                "tienda_codigo": col_upper,
                "tienda_nombre": tienda_nombre,
                "cantidad": cantidad_val,
                "validacion": validacion
            }
        nombre_tienda = col_name[len(self.utils.PREFIJO_TIENDA_ANTIGUO):].strip()
        return {
            "tienda_codigo": nombre_tienda.upper(),
            "tienda_nombre": nombre_tienda.title(),
            "cantidad": cantidad_val,
            "validacion": {"valida": True, "errores": [], "warnings": []}
        }

    def extraer_distribuciones_tienda(
        self,
        row: pd.Series,
        columnas: List[str],
        catalogo_tiendas: Optional[Dict[str, Dict[str, Any]]] = None
    ) -> List[Dict[str, Any]]:
        """
        Extrae distribuciones por tienda desde las columnas del Excel.

        Args:
            row: Fila de datos con las cantidades
            columnas: Lista de nombres de columnas
            catalogo_tiendas: Diccionario opcional con tiendas válidas

        Returns:
            Lista de distribuciones
        """
        distribuciones = []

        for col_name in columnas:
            col_lower = normalize_column_name(col_name)
            col_upper = str(col_name).upper().strip()

            es_codigo_tienda = bool(re.match(self.utils.PATRON_CODIGO_TIENDA, col_upper))
            es_formato_antiguo = col_lower.startswith(self.utils.PREFIJO_TIENDA_ANTIGUO)

            if not (es_codigo_tienda or es_formato_antiguo):
                continue

            cantidad = self.utils.get_numeric_value(row, col_name)
            if cantidad is None or cantidad < 0:
                continue

            cantidad_val = float(cantidad) if not isinstance(cantidad, float) else cantidad
            distribuciones.append(
                self._construir_distribucion_columna(col_name, col_upper, es_codigo_tienda, cantidad_val, catalogo_tiendas)
            )

        return distribuciones

    @staticmethod
    def _normalizar_encabezado_tiendas(text: str) -> str:
        """Normaliza un encabezado de columna eliminando tildes, espacios y guiones."""
        t = normalize_column_name(text)
        t = ''.join(c for c in unicodedata.normalize('NFD', t) if unicodedata.category(c) != 'Mn')
        return t.replace(' ', '').replace('-', '')

    @staticmethod
    def _encontrar_hoja_tiendas(workbook: Any) -> Optional[str]:
        """Retorna el nombre de la hoja 'Tiendas' (case-insensitive) o None si no existe."""
        for name in workbook.sheetnames:
            if name.lower() == "tiendas":
                return name
        return None

    @staticmethod
    def _leer_headers_tiendas(ws: Any) -> Dict[int, str]:
        """Construye un dict {col_idx: header_normalizado} de la primera fila."""
        headers: Dict[int, str] = {}
        for col_idx, cell in enumerate(ws[1]):
            if cell.value:
                headers[col_idx] = PedidoParserTiendas._normalizar_encabezado_tiendas(str(cell.value))
        return headers

    @staticmethod
    def _detectar_columnas_tiendas(headers: Dict[int, str]) -> tuple:
        """Detecta los índices de columna para código, nombre y activa."""
        codigo_candidates = {"idtienda", "codigo", "codigo_tienda", "codigotienda", "code"}
        nombre_candidates = {"nombretienda", "nombre", "tienda", "nombredetienda"}
        activa_candidates = {"activo", "activa", "active", "estado"}
        col_codigo = col_nombre = col_activa = None
        for col_idx, header in headers.items():
            if header in codigo_candidates or any(h in header for h in codigo_candidates):
                col_codigo = col_idx
            elif header in nombre_candidates or any(h in header for h in nombre_candidates):
                col_nombre = col_idx
            elif header in activa_candidates or any(h in header for h in activa_candidates):
                col_activa = col_idx
        return col_codigo, col_nombre, col_activa

    @staticmethod
    def _validar_sincronizacion_tiendas(ws: Any, col_codigo: int, col_nombre: int) -> None:
        """Lanza ValueError si alguna fila tiene código sin nombre o viceversa."""
        filas_desincronizadas = []
        for fila_idx in range(2, ws.max_row + 1):
            codigo_bruto = ws.cell(row=fila_idx, column=(col_codigo + 1)).value
            nombre_bruto = ws.cell(row=fila_idx, column=(col_nombre + 1)).value if col_nombre is not None else None
            has_codigo = bool(codigo_bruto is not None and str(codigo_bruto).strip())
            has_nombre = bool(nombre_bruto is not None and str(nombre_bruto).strip())
            if has_codigo != has_nombre:
                filas_desincronizadas.append(fila_idx)
        if filas_desincronizadas:
            filas_str = ", ".join(str(f) for f in filas_desincronizadas)
            raise ValueError(
                f"Error de sincronización entre ID TIENDA y NOMBRE TIENDA. "
                f"Filas problemáticas: {filas_str}. "
                f"Si hay información en la columna ID TIENDA debe haber en NOMBRE TIENDA y viceversa."
            )

    @staticmethod
    def _normalizar_activa_tienda(activa_bruto: Any) -> bool:
        """Convierte el valor crudo del campo activa a bool."""
        if activa_bruto is not None:
            activa_str = normalize_column_name(activa_bruto)
            if activa_str in ("no", "false", "0", "f"):
                return False
            elif activa_str in ("si", "yes", "true", "1", "v"):
                return True
        return True

    def _procesar_fila_tienda(
        self,
        ws: Any,
        fila_idx: int,
        col_codigo: int,
        col_nombre: int,
        col_activa: int,
    ) -> Optional[Dict[str, Any]]:
        """Procesa una fila de datos y retorna la entrada del catálogo o None si debe omitirse."""
        codigo_bruto = ws.cell(row=fila_idx, column=(col_codigo + 1)).value
        nombre_bruto = ws.cell(row=fila_idx, column=(col_nombre + 1)).value if col_nombre is not None else None
        activa_bruto = ws.cell(row=fila_idx, column=(col_activa + 1)).value if col_activa is not None else True
        has_codigo = bool(codigo_bruto is not None and str(codigo_bruto).strip())
        has_nombre = bool(nombre_bruto is not None and str(nombre_bruto).strip())
        if not has_codigo or not has_nombre:
            return None
        codigo_normalizado = str(codigo_bruto).strip().upper()
        if not re.match(self.utils.PATRON_CODIGO_VALIDO, codigo_normalizado):
            logger.warning(f"⚠️ Fila {fila_idx}: Código '{codigo_normalizado}' contiene caracteres inválidos")
            return None
        nombre_normalizado = str(nombre_bruto).strip() if nombre_bruto else codigo_normalizado
        return {
            "codigo": codigo_normalizado,
            "nombre": nombre_normalizado,
            "activa": PedidoParserTiendas._normalizar_activa_tienda(activa_bruto),
            "codigo_original": codigo_bruto,
            "fila_origen": fila_idx,
        }

    def leer_hoja_tiendas(self, file_content: bytes) -> Dict[str, Dict]:
        """
        Lee la hoja "Tiendas" y retorna catálogo NORMALIZADO.

        Args:
            file_content: Contenido del archivo Excel en bytes

        Returns:
            Dict con catálogo de tiendas normalizado
        """
        import openpyxl

        try:
            workbook = openpyxl.load_workbook(io.BytesIO(file_content), data_only=True)
            sheet_name = PedidoParserTiendas._encontrar_hoja_tiendas(workbook)
            if not sheet_name:
                logger.warning("⚠️ Hoja 'Tiendas' no encontrada en Excel")
                return {}
            ws = workbook[sheet_name]
            headers = PedidoParserTiendas._leer_headers_tiendas(ws)
            col_codigo, col_nombre, col_activa = PedidoParserTiendas._detectar_columnas_tiendas(headers)
            if col_codigo is None or col_nombre is None or col_activa is None:
                logger.warning("⚠️ Columnas requeridas no encontradas en hoja 'Tiendas'")
                return {}
            PedidoParserTiendas._validar_sincronizacion_tiendas(ws, col_codigo, col_nombre)
            catalogo: Dict[str, Dict] = {}
            for fila_idx in range(2, ws.max_row + 1):
                entrada = self._procesar_fila_tienda(ws, fila_idx, col_codigo, col_nombre, col_activa)
                if entrada:
                    catalogo[entrada["codigo"]] = entrada
            logger.info(f"✅ Catálogo de tiendas cargado: {len(catalogo)} tiendas")
            return catalogo
        except ValueError:
            raise
        except Exception as e:
            logger.error(f"❌ Error al leer hoja 'Tiendas': {e}", exc_info=True)
            return {}

"""Generador del Excel simple de Integración SAP (OC + Solicitudes de Transferencia).

Se descarga desde el modal "Integración con SAP Completada": la misma
información que se envía en el correo "OC y Solicitudes de Transferencia -
Listo para Despacho" (ver `src/utils/email_notifier.py::_build_despacho_email`),
en un formato mínimo de 2 hojas.
"""

from io import BytesIO
from typing import Any, Dict, Optional

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

# Orden y etiquetas exactas pedidas para la hoja de Solicitudes de Transferencia.
_HEADERS_TRANSFERENCIAS = ["DOCNUM", "NOMBRE TIENDA", "CODIGO TIENDA"]


def _hoja_orden_compra(wb: Workbook, oc_cabecera: Dict[str, Any], doc_num: Any, doc_entry: Any, pedido_numero: Any) -> None:
    ws = wb.create_sheet("Orden de Compra")
    numero_oc_sap = doc_num or oc_cabecera.get("numero_oc_sap") or "N/A"
    campos = [
        ("Número OC SAP", numero_oc_sap),
        ("DocEntry", doc_entry or "N/A"),
        ("Proveedor", oc_cabecera.get("proveedor") or "N/A"),
        ("Referencia", oc_cabecera.get("referencia_despacho") or "N/A"),
        ("Pedido", pedido_numero or "N/A"),
    ]
    for row, (campo, valor) in enumerate(campos, 1):
        ws.cell(row=row, column=1, value=campo).font = Font(bold=True)
        ws.cell(row=row, column=2, value=valor)
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 40


def _hoja_solicitudes_transferencia(wb: Workbook, solicitudes_transferencia: Optional[Dict[str, Any]]) -> None:
    ws = wb.create_sheet("Solicitudes de Transferencia")
    for col, header in enumerate(_HEADERS_TRANSFERENCIAS, 1):
        ws.cell(row=1, column=col, value=header).font = Font(bold=True)

    resultados = (solicitudes_transferencia or {}).get("resultados") or []
    for row, resultado in enumerate(resultados, 2):
        ws.cell(row=row, column=1, value=resultado.get("doc_num"))
        ws.cell(row=row, column=2, value=resultado.get("tienda_nombre") or resultado.get("tienda_codigo"))
        ws.cell(row=row, column=3, value=resultado.get("tienda_codigo"))

    for col in range(1, len(_HEADERS_TRANSFERENCIAS) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 22


def generar_excel_integracion_sap(
    oc_cabecera: Dict[str, Any],
    doc_num: Any,
    doc_entry: Any,
    solicitudes_transferencia: Optional[Dict[str, Any]],
    pedido_numero: Any,
) -> BytesIO:
    """Genera el Excel de 2 hojas de integración SAP para descarga desde el modal.

    Hoja 1 "Orden de Compra": Número OC SAP, DocEntry, Proveedor, Referencia y
    Pedido (mismos campos que el correo de despacho).
    Hoja 2 "Solicitudes de Transferencia": una fila por resultado, con
    exclusivamente DOCNUM / NOMBRE TIENDA / CODIGO TIENDA.

    Returns:
        BytesIO posicionado en 0, listo para StreamingResponse.
    """
    wb = Workbook()
    default_sheet = wb.active
    if default_sheet is not None:
        wb.remove(default_sheet)

    _hoja_orden_compra(wb, oc_cabecera, doc_num, doc_entry, pedido_numero)
    _hoja_solicitudes_transferencia(wb, solicitudes_transferencia)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer

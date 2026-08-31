"""Generadores de archivos Excel para redistribución."""

from io import BytesIO
from typing import Any, Dict

import pandas as pd
from openpyxl.styles import Border, Font, PatternFill, Side


def generar_excel_reporte_diferencias(
    diferencias: Dict[str, Any],
) -> BytesIO:
    """Genera el libro Excel de reporte de diferencias."""
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        resumen_data = {
            "Concepto": [
                "Total Productos",
                "Total Cantidad Solicitada",
                "Total Cantidad Confirmada",
                "Total Cantidad Asignada",
                "Total Diferencia Cantidad",
                "Total Costo Original",
                "Total Costo Ajustado",
                "Total Diferencia Costo",
            ],
            "Valor": [
                diferencias.get("total_productos"),
                diferencias.get("total_cantidad_solicitada"),
                diferencias.get("total_cantidad_confirmada"),
                diferencias.get("total_cantidad_asignada"),
                diferencias.get("total_diferencia_cantidad"),
                diferencias.get("total_costo_original") or 0,
                diferencias.get("total_costo_ajustado") or 0,
                diferencias.get("total_diferencia_costo") or 0,
            ],
        }
        pd.DataFrame(resumen_data).to_excel(writer, sheet_name="Resumen", index=False)

        if diferencias.get("diferencias_productos"):
            df_productos = pd.DataFrame(diferencias["diferencias_productos"])
            df_productos.to_excel(writer, sheet_name="Diferencias por Producto", index=False)

        if diferencias.get("diferencias_tiendas"):
            df_tiendas = pd.DataFrame(diferencias["diferencias_tiendas"])
            df_tiendas.to_excel(writer, sheet_name="Diferencias por Tienda", index=False)

    output.seek(0)
    return output


def _crear_estilos() -> Dict[str, Any]:
    """Compartido con _excel_traslados.py (generar_excel_traslados)."""
    return {
        "header_fill": PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid"),
        "header_font": Font(bold=True, color="FFFFFF", size=12),
        "title_fill": PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid"),
        "title_font": Font(bold=True, size=14),
        "border": Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        ),
    }

"""Generador del Excel de Solicitudes de Traslado (requerimiento 1.9).

Exporta, por cada par (producto, tienda destino) de la distribución confirmada:
número de traslado, nombre de tienda, código de barras, cantidad solicitada y
cantidad confirmada.

Notas de alcance verificadas contra el esquema y los datos:

- El Nº de traslado se lee SOLO de base de datos (solicitud_transferencia_detalle),
  nunca de una consulta en vivo a SAP, para que el archivo no dependa de que el
  middleware esté disponible.
- El Nº de traslado existe por TIENDA destino, no por producto: el sistema no
  registra qué producto viajó en qué ITR, así que el mismo número se repite en
  todas las líneas de una tienda.
- Las BODEGAS centrales (RL-CD, RG-CD, OUT-CD) se incluyen para que los totales
  cuadren con la redistribución real, pero su flujo de ITR no persiste el doc_num,
  así que su Nº de traslado sale como "No aplica".
- Solo se leen cantidades. Los campos de costo/subtotal de la distribución tienen
  un desfase de índice conocido en el repositorio y quedan deliberadamente fuera.
"""

from io import BytesIO
from typing import Any, Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from ._excel_generators import _crear_estilos

# Texto usado cuando la tienda destino es una bodega central, cuyo flujo de ITR
# no persiste el doc_num en base de datos.
SIN_TRASLADO_BODEGA = "No aplica"

# Texto usado cuando la tienda sí podría tener Nº de traslado pero aún no existe
# (la ODC todavía no se envió a SAP, o la ITR falló y no se ha reintentado).
SIN_TRASLADO_PENDIENTE = "Pendiente"


def _norm(codigo: Optional[str]) -> str:
    """Normaliza un código de tienda para el lookup del mapa de ITR."""
    return (codigo or "").strip().upper()


def _a_float(valor: Any) -> float:
    """Convierte a float tolerando None, Decimal y cadenas vacías."""
    if valor is None or valor == "":
        return 0.0
    try:
        return float(valor)
    except (TypeError, ValueError):
        return 0.0


def _doc_num_de(itr_por_tienda: Dict[str, int], tienda_codigo: str) -> Optional[int]:
    """Resuelve el doc_num de la ITR de una tienda, o None si no hay."""
    return itr_por_tienda.get(_norm(tienda_codigo))


def construir_matriz_traslados(
    productos_enriquecidos: List[Dict[str, Any]],
    itr_por_tienda: Optional[Dict[str, int]] = None,
) -> Dict[str, Any]:
    """Construye la estructura de datos del export a partir de la distribución.

    Función pura (sin openpyxl) para poder probarla con diccionarios.

    Args:
        productos_enriquecidos: productos de la distribución ya enriquecidos con
            codigo_interno, descripcion_sap y codigo_barras del pedido.
        itr_por_tienda: mapa tienda_codigo NORMALIZADO -> doc_num de la ITR.

    Returns:
        Dict con las tiendas del eje, las filas (producto, tienda), los totales
        por tienda y la lista de tiendas sin Nº de traslado.
    """
    itr_por_tienda = itr_por_tienda or {}

    tiendas_por_codigo: Dict[str, Dict[str, Any]] = {}
    filas: List[Dict[str, Any]] = []
    totales_por_tienda: Dict[str, Dict[str, float]] = {}

    for producto in productos_enriquecidos:
        codigo_articulo = producto.get("codigo_interno") or ""
        descripcion = producto.get("descripcion_sap") or ""
        codigo_barras = producto.get("codigo_barras") or ""
        # Confirmado real del proveedor, que solo existe a nivel producto.
        confirmado_producto = _a_float(producto.get("cantidad_confirmada_total"))

        for dist in producto.get("distribuciones") or []:
            tienda_codigo = dist.get("tienda_codigo") or ""
            if not tienda_codigo:
                continue

            es_bodega = bool(dist.get("es_bodega"))
            if tienda_codigo not in tiendas_por_codigo:
                tiendas_por_codigo[tienda_codigo] = {
                    "codigo": tienda_codigo,
                    "nombre": dist.get("tienda_nombre") or "",
                    "es_bodega": es_bodega,
                    "doc_num": _doc_num_de(itr_por_tienda, tienda_codigo),
                }
                totales_por_tienda[tienda_codigo] = {
                    "solicitada": 0.0,
                    "confirmada": 0.0,
                    "asignada": 0.0,
                }

            solicitada = _a_float(dist.get("cantidad_solicitada"))
            confirmada = _a_float(dist.get("cantidad_confirmada"))
            asignada = _a_float(dist.get("cantidad_asignada"))

            filas.append(
                {
                    "doc_num": tiendas_por_codigo[tienda_codigo]["doc_num"],
                    "tienda_codigo": tienda_codigo,
                    "tienda_nombre": tiendas_por_codigo[tienda_codigo]["nombre"],
                    "es_bodega": es_bodega,
                    "codigo_articulo": codigo_articulo,
                    "descripcion": descripcion,
                    "codigo_barras": codigo_barras,
                    "cantidad_solicitada": solicitada,
                    "cantidad_confirmada": confirmada,
                    "cantidad_asignada": asignada,
                    "confirmado_total_producto": confirmado_producto,
                }
            )

            totales_por_tienda[tienda_codigo]["solicitada"] += solicitada
            totales_por_tienda[tienda_codigo]["confirmada"] += confirmada
            totales_por_tienda[tienda_codigo]["asignada"] += asignada

    # Tiendas primero (ordenadas por código), bodegas al final.
    tiendas = sorted(
        tiendas_por_codigo.values(),
        key=lambda t: (t["es_bodega"], t["codigo"]),
    )

    # Las filas siguen el mismo criterio de orden que el eje de tiendas, para que
    # el archivo se lea igual que la pantalla de redistribución.
    orden_tienda = {t["codigo"]: idx for idx, t in enumerate(tiendas)}
    filas.sort(
        key=lambda f: (orden_tienda.get(f["tienda_codigo"], 0), f["codigo_articulo"])
    )

    return {
        "tiendas": tiendas,
        "filas": filas,
        "totales_por_tienda": totales_por_tienda,
        "gran_total_solicitada": sum(f["cantidad_solicitada"] for f in filas),
        "gran_total_confirmada": sum(f["cantidad_confirmada"] for f in filas),
        "gran_total_asignada": sum(f["cantidad_asignada"] for f in filas),
        "tiendas_sin_doc_num": [
            t["codigo"] for t in tiendas if t["doc_num"] is None and not t["es_bodega"]
        ],
        "bodegas": [t["codigo"] for t in tiendas if t["es_bodega"]],
    }


def agrupar_filas_por_producto(filas: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """Pivotea las filas planas (producto, tienda) a una fila por producto.

    Función pura, independiente de openpyxl, para armar la matriz con tiendas en
    columnas. Ni la cantidad solicitada ni la confirmada se muestran aquí: la
    matriz refleja la redistribución (cantidad_asignada), la misma que se ve
    por tienda en la pantalla de "Ajuste y Redistribución a Tiendas".

    Args:
        filas: lista de filas de `construir_matriz_traslados()["filas"]`.

    Returns:
        Lista de productos en el orden de primera aparición, cada uno con
        `celdas` (cantidad_asignada, indexada por tienda_codigo) y
        `total_asignada` (suma de la redistribución en todas las tiendas).
    """
    productos: Dict[str, Dict[str, Any]] = {}
    orden: List[str] = []

    for fila in filas:
        codigo = fila["codigo_articulo"]
        if codigo not in productos:
            productos[codigo] = {
                "codigo_articulo": codigo,
                "descripcion": fila["descripcion"],
                "codigo_barras": fila["codigo_barras"],
                "celdas": {},
                "total_asignada": 0.0,
            }
            orden.append(codigo)

        producto = productos[codigo]
        producto["celdas"][fila["tienda_codigo"]] = fila["cantidad_asignada"]
        producto["total_asignada"] += fila["cantidad_asignada"]

    return [productos[codigo] for codigo in orden]


def _texto_traslado(fila: Dict[str, Any]) -> Any:
    """Devuelve el valor de la columna Nº de Traslado para una fila."""
    if fila["doc_num"] is not None:
        return fila["doc_num"]
    return SIN_TRASLADO_BODEGA if fila["es_bodega"] else SIN_TRASLADO_PENDIENTE


def _hoja_cabecera_traslados(
    wb,
    pedido_id: str,
    cabecera: Dict[str, Any],
    distribucion_data: Dict[str, Any],
    matriz: Dict[str, Any],
    user: Dict[str, Any],
    estilos: Dict[str, Any],
):
    """Hoja de contexto: identifica el pedido y explicita el alcance del archivo."""
    from src.utils.timezone import ahora_naive

    ws = wb.create_sheet("Información del Pedido")
    ws["A1"] = "SOLICITUDES DE TRASLADO - INFORMACIÓN DEL PEDIDO"
    ws["A1"].font = estilos["title_font"]
    ws["A1"].fill = estilos["title_fill"]
    ws.merge_cells("A1:B1")

    tiendas_sin = matriz.get("tiendas_sin_doc_num") or []
    bodegas = matriz.get("bodegas") or []
    con_doc_num = sum(1 for t in matriz.get("tiendas", []) if t["doc_num"] is not None)

    info_fields = [
        ("Número de Pedido", cabecera.get("numero_pedido", "N/A")),
        ("Pedido ID", pedido_id),
        ("Versión Redistribución", distribucion_data.get("version_actual", "N/A")),
        ("Estado", cabecera.get("estado_nombre", "N/A")),
        ("Fecha Pedido", cabecera.get("fecha_pedido", "N/A")),
        ("Fecha Entrega", cabecera.get("fecha_entrega", "N/A")),
        ("Origen del Nº de traslado", "Base de datos (solicitud_transferencia_detalle)"),
        ("Tiendas con Nº de traslado", con_doc_num),
        ("Tiendas sin Nº de traslado", ", ".join(tiendas_sin) if tiendas_sin else "Ninguna"),
        ("Bodegas incluidas", ", ".join(bodegas) if bodegas else "Ninguna"),
        (
            "Alcance del Nº de traslado",
            "Por tienda destino. El sistema no registra qué producto viajó en qué "
            "solicitud de transferencia, por lo que el número se repite en todas las "
            "líneas de la misma tienda. Las bodegas centrales no tienen número "
            "registrado en base de datos.",
        ),
        (
            "Cantidad Confirmada",
            "La hoja 'Matriz por Tienda' no la muestra: por tienda y en 'Total "
            "General' se ve la redistribución (Cantidad Asignada), igual que en la "
            "pantalla de 'Ajuste y Redistribución a Tiendas'. En la hoja de detalle "
            "sí se muestra por tienda (distribución proporcional del confirmado "
            "real), y el confirmado global del proveedor aparece aparte en la "
            "columna 'Confirmado Total Producto'.",
        ),
        (
            "Cantidad Asignada (ITR)",
            "Cantidad que efectivamente viaja en la Solicitud de Transferencia a SAP.",
        ),
    ]

    row = 3
    for field, value in info_fields:
        ws[f"A{row}"] = field
        ws[f"A{row}"].font = Font(bold=True)
        ws[f"B{row}"] = value
        ws[f"B{row}"].alignment = Alignment(wrap_text=True, vertical="top")
        row += 1

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 80

    row = ws.max_row + 2
    ws[f"A{row}"] = "Descargado por:"
    ws[f"A{row}"].font = Font(bold=True)
    ws[f"B{row}"] = user.get("username", "Desconocido")
    row += 1
    ws[f"A{row}"] = "Fecha de descarga:"
    ws[f"A{row}"].font = Font(bold=True)
    ws[f"B{row}"] = ahora_naive().strftime("%Y-%m-%d %H:%M:%S")


# Encabezados de la hoja de detalle. Las cinco primeras columnas del requerimiento
# son A (Número de Traslado), C (Nombre de la Tienda), F (Código de Barras),
# G (Cantidad Solicitada) y H (Cantidad Confirmada).
_HEADERS_TRASLADOS = [
    "Número de Traslado",
    "Código Tienda",
    "Nombre de la Tienda",
    "Código Artículo",
    "Descripción",
    "Código de Barras",
    "Cantidad Solicitada",
    "Cantidad Confirmada",
    "Cantidad Asignada (ITR)",
    "Confirmado Total Producto",
]

_ANCHOS_TRASLADOS = [22, 15, 32, 18, 45, 20, 18, 19, 21, 24]


def _hoja_traslados_detalle(wb, matriz: Dict[str, Any], estilos: Dict[str, Any]):
    """Hoja principal: una fila por (producto, tienda destino)."""
    ws = wb.create_sheet("Traslados (Detalle)")
    ws["A1"] = "SOLICITUDES DE TRASLADO - DETALLE POR TIENDA"
    ws["A1"].font = estilos["title_font"]
    ws["A1"].fill = estilos["title_fill"]
    ws.merge_cells(f"A1:{chr(ord('A') + len(_HEADERS_TRASLADOS) - 1)}1")

    for col_num, header in enumerate(_HEADERS_TRASLADOS, 1):
        cell = ws.cell(row=3, column=col_num)
        cell.value = header
        cell.font = estilos["header_font"]
        cell.fill = estilos["header_fill"]
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = estilos["border"]

    row = 4
    for fila in matriz.get("filas", []):
        ws.cell(row=row, column=1, value=_texto_traslado(fila))
        ws.cell(row=row, column=2, value=fila["tienda_codigo"])
        ws.cell(row=row, column=3, value=fila["tienda_nombre"])
        ws.cell(row=row, column=4, value=fila["codigo_articulo"])
        ws.cell(row=row, column=5, value=fila["descripcion"])
        ws.cell(row=row, column=6, value=fila["codigo_barras"])
        ws.cell(row=row, column=7, value=fila["cantidad_solicitada"])
        ws.cell(row=row, column=8, value=fila["cantidad_confirmada"])
        ws.cell(row=row, column=9, value=fila["cantidad_asignada"])
        ws.cell(row=row, column=10, value=fila["confirmado_total_producto"])
        row += 1

    # Fila de totales, separada por una fila en blanco (igual que _hoja_resumen).
    row += 1
    ws.cell(row=row, column=1, value="TOTALES").font = Font(bold=True)
    ws.cell(row=row, column=7, value=matriz.get("gran_total_solicitada", 0.0)).font = Font(bold=True)
    ws.cell(row=row, column=8, value=matriz.get("gran_total_confirmada", 0.0)).font = Font(bold=True)
    ws.cell(row=row, column=9, value=matriz.get("gran_total_asignada", 0.0)).font = Font(bold=True)

    for idx, ancho in enumerate(_ANCHOS_TRASLADOS, 1):
        ws.column_dimensions[chr(ord("A") + idx - 1)].width = ancho

    # Encabezado fijo al desplazarse (mismo criterio de usabilidad del punto 1.1.1).
    ws.freeze_panes = "A4"


_COL_CODIGO_ARTICULO = 1
_COL_CODIGO_BARRAS = 2
_COL_FIJAS = 2  # cantidad de columnas fijas antes del bloque de tiendas
_FILA_DOCNUM = 3
_FILA_TIENDA = 4
_FILA_SUBHEADER = 5
_FILA_DATOS_INICIO = 6
_ANCHO_COL_TIENDA = 9


def _pintar_encabezado_bloque(ws, fila: int, col_inicio: int, col_fin: int, valor: Any, estilos: Dict[str, Any]):
    """Escribe y da estilo de encabezado a un bloque de celdas fusionadas de 2 columnas."""
    ws.merge_cells(start_row=fila, start_column=col_inicio, end_row=fila, end_column=col_fin)
    for col in range(col_inicio, col_fin + 1):
        cell = ws.cell(row=fila, column=col)
        cell.fill = estilos["header_fill"]
        cell.font = estilos["header_font"]
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = estilos["border"]
    ws.cell(row=fila, column=col_inicio, value=valor)


def _pintar_encabezado_tienda(ws, col: int, tienda: Dict[str, Any], estilos: Dict[str, Any]) -> None:
    """Escribe el bloque de encabezado (1 columna) de una tienda: DocNum + nombre.

    El nombre de la tienda ocupa la fila de nombre y la de subencabezado
    fusionadas (una sola columna por tienda, sin subencabezado adicional: el
    dato que se muestra ahí es la redistribución, cantidad_asignada).
    """
    if tienda["doc_num"] is not None:
        valor_doc_num: Any = tienda["doc_num"]
    else:
        valor_doc_num = SIN_TRASLADO_BODEGA if tienda["es_bodega"] else SIN_TRASLADO_PENDIENTE

    docnum_cell = ws.cell(row=_FILA_DOCNUM, column=col, value=valor_doc_num)
    docnum_cell.fill = estilos["header_fill"]
    docnum_cell.font = estilos["header_font"]
    docnum_cell.alignment = Alignment(horizontal="center", vertical="center", text_rotation=90, wrap_text=True)
    docnum_cell.border = estilos["border"]
    if isinstance(valor_doc_num, int):
        # Sin esto, Excel puede mostrar el DocNum en notación científica (ej.
        # "7E+07") cuando la columna es angosta y hay muchas tiendas.
        docnum_cell.number_format = "0"

    ws.merge_cells(start_row=_FILA_TIENDA, start_column=col, end_row=_FILA_SUBHEADER, end_column=col)
    nombre_cell = ws.cell(row=_FILA_TIENDA, column=col, value=tienda["nombre"])
    nombre_cell.font = Font(bold=True)
    nombre_cell.alignment = Alignment(horizontal="center", vertical="center", text_rotation=90, wrap_text=True)
    nombre_cell.border = estilos["border"]
    ws.cell(row=_FILA_SUBHEADER, column=col).border = estilos["border"]


def _hoja_matriz_traslados(wb, pedido_data: Dict[str, Any], matriz: Dict[str, Any], estilos: Dict[str, Any]):
    """Hoja principal: tiendas en columnas (1 por tienda: cantidad_asignada).

    Ni lo solicitado ni lo confirmado se muestran aquí: la cantidad que se ve
    por tienda es la de la redistribución (cantidad_asignada), la misma que
    la pantalla "Ajuste y Redistribución a Tiendas" muestra por tienda.
    "Total General" es una sola columna con la suma de esa redistribución.
    """
    ws = wb.create_sheet("Matriz por Tienda")
    numero_pedido = pedido_data.get("cabecera", {}).get("numero_pedido") or "SOLICITUDES DE TRASLADO"

    ws.cell(row=1, column=1, value=numero_pedido)
    ws.cell(row=1, column=1).font = estilos["title_font"]
    ws.cell(row=1, column=1).fill = estilos["title_fill"]
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2)

    tiendas = matriz.get("tiendas", [])
    productos = agrupar_filas_por_producto(matriz.get("filas", []))

    # Columnas fijas (Código Artículo / Código de Barras), fusionadas verticalmente
    # para alinear con las 3 filas de encabezado del bloque de tiendas.
    for col, etiqueta in ((_COL_CODIGO_ARTICULO, "Código Artículo"), (_COL_CODIGO_BARRAS, "Código de Barras")):
        ws.merge_cells(start_row=_FILA_DOCNUM, start_column=col, end_row=_FILA_SUBHEADER, end_column=col)
        cell = ws.cell(row=_FILA_DOCNUM, column=col, value=etiqueta)
        cell.font = estilos["header_font"]
        cell.fill = estilos["header_fill"]
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = estilos["border"]

    col_tienda: Dict[str, int] = {}
    col = _COL_FIJAS + 1
    for tienda in tiendas:
        col_tienda[tienda["codigo"]] = col
        _pintar_encabezado_tienda(ws, col, tienda, estilos)
        col += 1

    col_total = col
    # "Total General" ocupa una sola columna, fusionada en las 3 filas de
    # encabezado (no tiene DocNum propio ni subencabezado que mostrar).
    ws.merge_cells(
        start_row=_FILA_DOCNUM, start_column=col_total,
        end_row=_FILA_SUBHEADER, end_column=col_total,
    )
    cell = ws.cell(row=_FILA_DOCNUM, column=col_total, value="Total General")
    cell.fill = estilos["header_fill"]
    cell.font = estilos["header_font"]
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = estilos["border"]

    fila = _FILA_DATOS_INICIO
    for producto in productos:
        ws.cell(row=fila, column=_COL_CODIGO_ARTICULO, value=producto["codigo_articulo"])
        ws.cell(row=fila, column=_COL_CODIGO_BARRAS, value=producto["codigo_barras"])

        for tienda in tiendas:
            asignada = producto["celdas"].get(tienda["codigo"], 0.0)
            ws.cell(row=fila, column=col_tienda[tienda["codigo"]], value=asignada)

        total_cell = ws.cell(row=fila, column=col_total, value=producto["total_asignada"])
        total_cell.font = Font(bold=True)
        fila += 1

    ws.cell(row=fila, column=_COL_CODIGO_ARTICULO, value="TOTALES").font = Font(bold=True)
    totales_por_tienda = matriz.get("totales_por_tienda", {})
    for tienda in tiendas:
        acumulado = totales_por_tienda.get(tienda["codigo"], {})
        ws.cell(row=fila, column=col_tienda[tienda["codigo"]], value=acumulado.get("asignada", 0.0)).font = Font(bold=True)

    ws.cell(row=fila, column=col_total, value=matriz.get("gran_total_asignada", 0.0)).font = Font(bold=True)

    ws.column_dimensions[get_column_letter(_COL_CODIGO_ARTICULO)].width = 20
    ws.column_dimensions[get_column_letter(_COL_CODIGO_BARRAS)].width = 20
    for c in range(_COL_FIJAS + 1, col_total + 1):
        ws.column_dimensions[get_column_letter(c)].width = _ANCHO_COL_TIENDA
    # Suficiente para que un DocNum de 8 dígitos (o "No aplica") quepa en una
    # sola línea rotada 90°: en vertical, cada carácter consume espacio de
    # ALTURA de fila (no de ancho de columna), y los números nunca hacen wrap
    # (a diferencia del texto): si no caben, Excel los muestra como "#######"
    # en vez de partir la línea.
    ws.row_dimensions[_FILA_DOCNUM].height = 100
    ws.row_dimensions[_FILA_TIENDA].height = 90

    # Encabezado (3 filas) y columnas fijas congeladas al desplazarse.
    ws.freeze_panes = ws.cell(row=_FILA_DATOS_INICIO, column=_COL_FIJAS + 1).coordinate


def generar_excel_traslados(
    pedido_id: str,
    pedido_data: Dict[str, Any],
    distribucion_data: Dict[str, Any],
    productos_enriquecidos: List[Dict[str, Any]],
    user: Dict[str, Any],
    itr_por_tienda: Optional[Dict[str, int]] = None,
) -> BytesIO:
    """Genera el libro Excel de Solicitudes de Traslado.

    Args:
        pedido_id: UUID del pedido.
        pedido_data: pedido completo (cabecera + detalles).
        distribucion_data: distribución confirmada (incluye version_actual).
        productos_enriquecidos: productos con codigo_interno/descripcion_sap/codigo_barras.
        user: usuario autenticado, para la trazabilidad de la descarga.
        itr_por_tienda: mapa tienda_codigo normalizado -> doc_num de la ITR.

    Returns:
        BytesIO posicionado en 0, listo para StreamingResponse.
    """
    cabecera = pedido_data.get("cabecera", {})
    matriz = construir_matriz_traslados(productos_enriquecidos, itr_por_tienda)

    wb = Workbook()
    default_sheet = wb.active
    if default_sheet is not None:
        wb.remove(default_sheet)

    estilos = _crear_estilos()
    _hoja_cabecera_traslados(
        wb, pedido_id, cabecera, distribucion_data, matriz, user, estilos
    )
    _hoja_matriz_traslados(wb, pedido_data, matriz, estilos)
    _hoja_traslados_detalle(wb, matriz, estilos)

    excel_buffer = BytesIO()
    wb.save(excel_buffer)
    excel_buffer.seek(0)
    return excel_buffer

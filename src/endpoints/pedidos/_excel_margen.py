"""Generador del Excel de Análisis de Margen (requerimiento 1.4/1.5, SIS-PYTD-0002).

Exporta, por producto del pedido: código, descripción, PVP, Costo (columna "Costo"
de la plantilla Excel, informativa, req. 1.4), Costo Unitario (costo_doc_sin_iva /
unidad_conversion, insumo real del Margen, ver pedido_utils.calcular_margen_rentabilidad)
y el Margen de Rentabilidad calculado.

No incluye costo_doc_sin_iva (el costo/precio negociado que alimenta la OC): este
archivo es exclusivamente el análisis de margen, no un reemplazo de "Descargar Pedido".
"""

from io import BytesIO
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from .redistribucion._excel_generators import _crear_estilos

_HEADERS_MARGEN = [
    "Código",
    "Descripción",
    "Código de Barras",
    "Pack",
    "PVP",
    "Costo (Excel)",
    "Costo Unitario",
    "Margen de Rentabilidad",
    "Total",
]

_ANCHOS_MARGEN = [16, 45, 20, 10, 14, 14, 14, 20, 12]

_FORMATO_MONEDA = "#,##0.00"


def _hoja_cabecera_margen(
    wb,
    pedido_id: str,
    cabecera: Dict[str, Any],
    user: Dict[str, Any],
    estilos: Dict[str, Any],
):
    """Hoja de contexto: identifica el pedido y quién/cuándo descargó el análisis."""
    from src.utils.timezone import ahora_naive

    ws = wb.create_sheet("Información del Pedido")
    ws["A1"] = "ANÁLISIS DE MARGEN - INFORMACIÓN DEL PEDIDO"
    ws["A1"].font = estilos["title_font"]
    ws["A1"].fill = estilos["title_fill"]
    ws.merge_cells("A1:B1")

    info_fields = [
        ("Número de Pedido", cabecera.get("numero_pedido", "N/A")),
        ("Pedido ID", pedido_id),
        ("Estado", cabecera.get("estado_nombre", "N/A")),
        ("Fecha Pedido", cabecera.get("fecha_pedido", "N/A")),
        ("Fecha Entrega", cabecera.get("fecha_entrega", "N/A")),
        (
            "Fórmula del Margen",
            "1 - (Costo Unitario / (PVP / 1.15))",
        ),
        (
            "Costo Unitario",
            "costo_doc_sin_iva / unidad_conversion (pack ya resuelto SAP/Excel); "
            "es el insumo real del cálculo del Margen.",
        ),
        (
            "Costo (Excel)",
            "Columna 'Costo' cargada en la plantilla (req. 1.4): únicamente "
            "informativa, no alimenta el cálculo del Margen ni SAP (RN-05).",
        ),
    ]

    row = 3
    for field, value in info_fields:
        ws[f"A{row}"] = field
        ws[f"A{row}"].font = Font(bold=True)
        ws[f"B{row}"] = value
        ws[f"B{row}"].alignment = Alignment(wrap_text=True, vertical="top")
        row += 1

    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 80

    row = ws.max_row + 2
    ws[f"A{row}"] = "Descargado por:"
    ws[f"A{row}"].font = Font(bold=True)
    ws[f"B{row}"] = user.get("username", "Desconocido")
    row += 1
    ws[f"A{row}"] = "Fecha de descarga:"
    ws[f"A{row}"].font = Font(bold=True)
    ws[f"B{row}"] = ahora_naive().strftime("%Y-%m-%d %H:%M:%S")


def _hoja_margen_detalle(wb, detalles: List[Dict[str, Any]], estilos: Dict[str, Any]):
    """Hoja principal: una fila por producto del pedido."""
    ws = wb.create_sheet("Análisis de Margen")
    ws["A1"] = "ANÁLISIS DE MARGEN POR PRODUCTO"
    ws["A1"].font = estilos["title_font"]
    ws["A1"].fill = estilos["title_fill"]
    ws.merge_cells(f"A1:{get_column_letter(len(_HEADERS_MARGEN))}1")

    for col_num, header in enumerate(_HEADERS_MARGEN, 1):
        cell = ws.cell(row=3, column=col_num)
        cell.value = header
        cell.font = estilos["header_font"]
        cell.fill = estilos["header_fill"]
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = estilos["border"]

    row = 4
    total_cantidad = 0.0
    for detalle in detalles:
        ws.cell(row=row, column=1, value=detalle.get("codigo_interno"))
        ws.cell(row=row, column=2, value=detalle.get("descripcion_sap"))
        ws.cell(row=row, column=3, value=detalle.get("codigo_barras"))
        ws.cell(row=row, column=4, value=detalle.get("pack"))

        pvp_cell = ws.cell(row=row, column=5, value=detalle.get("pvp"))
        pvp_cell.number_format = _FORMATO_MONEDA

        costo_excel_cell = ws.cell(row=row, column=6, value=detalle.get("costo_margen"))
        costo_excel_cell.number_format = _FORMATO_MONEDA

        costo_unitario_cell = ws.cell(row=row, column=7, value=detalle.get("costo_unitario"))
        costo_unitario_cell.number_format = _FORMATO_MONEDA

        margen = detalle.get("margen")
        margen_cell = ws.cell(row=row, column=8, value=margen)
        margen_cell.number_format = "0.00%"
        if isinstance(margen, (int, float)) and margen < 0:
            margen_cell.font = Font(color="C00000")

        total = detalle.get("total") or 0.0
        ws.cell(row=row, column=9, value=total)
        total_cantidad += total

        for col in range(1, len(_HEADERS_MARGEN) + 1):
            ws.cell(row=row, column=col).border = estilos["border"]

        row += 1

    row += 1
    ws.cell(row=row, column=1, value="TOTALES").font = Font(bold=True)
    ws.cell(row=row, column=9, value=total_cantidad).font = Font(bold=True)

    for idx, ancho in enumerate(_ANCHOS_MARGEN, 1):
        ws.column_dimensions[get_column_letter(idx)].width = ancho

    # Encabezado fijo al desplazarse (mismo criterio de usabilidad del punto 1.1.1).
    ws.freeze_panes = "A4"


def generar_excel_margen(pedido_id: str, pedido_data: Dict[str, Any], user: Dict[str, Any]) -> BytesIO:
    """Genera el libro Excel del Análisis de Margen de un pedido (req. 1.4/1.5).

    Args:
        pedido_id: UUID del pedido.
        pedido_data: pedido completo (cabecera + detalles), ya con costo_margen,
            costo_unitario y margen calculados por PedidoRepository.
        user: usuario autenticado, para la trazabilidad de la descarga.

    Returns:
        BytesIO posicionado en 0, listo para servirse como archivo descargable.
    """
    cabecera = pedido_data.get("cabecera", {})
    detalles = pedido_data.get("detalles", [])

    wb = Workbook()
    default_sheet = wb.active
    if default_sheet is not None:
        wb.remove(default_sheet)

    estilos = _crear_estilos()
    _hoja_cabecera_margen(wb, pedido_id, cabecera, user, estilos)
    _hoja_margen_detalle(wb, detalles, estilos)

    excel_buffer = BytesIO()
    wb.save(excel_buffer)
    excel_buffer.seek(0)
    return excel_buffer

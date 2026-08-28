"""
Tests para la columna "Costo" (req. 1.4/1.5, SIS-PYTD-0002) en el parser de Excel.

Cubre:
- "Costo sin IVA" sigue mapeando a costo_doc_sin_iva (RN-05: no se toca ese campo).
- "Costo" (columna nueva) mapea a costo_margen, con valor distinto e independiente.
- La columna "Costo"/"Costo Unitario" no interfiere con la detección de bodegas
  (RL-CD/OUT-CD/RG-CD) ni de tiendas sin importar en qué posición de la hoja "detalle"
  se inserte (al final, o en medio como en la plantilla real "Costo Unitario" en H).

NOTA: la detección de bodegas/tiendas ya NO depende de una posición fija de columna
(ver identificar_columnas_bodegas/identificar_columnas_tiendas en pedido_parser_detalle.py):
excluye las columnas ya identificadas por nombre como campos de producto y evalúa el
resto por nombre/código. Antes dependía de INDICES_BODEGAS/INDICE_INICIO_TIENDAS
(constantes fijas, eliminadas), lo que se rompía silenciosamente en RG-CD al insertar
"Costo Unitario" antes de las columnas de bodega (incidente 2026-08-07).
"""

import io
from typing import Any, Dict, List, Optional

import pandas as pd
from openpyxl import Workbook

from src.database.services.pedido_parser.pedido_parser_detalle import PedidoParserDetalle


def _crear_excel_detalle(headers: List[str], filas: List[List[Any]]) -> bytes:
    """Crea un .xlsx mínimo con hoja 'detalle': fila 1 nombres, fila 2 códigos, fila 3+ datos."""
    wb = Workbook()
    ws = wb.active
    ws.title = "detalle"
    for col_idx, header in enumerate(headers, start=1):
        ws.cell(row=1, column=col_idx, value=header)
        ws.cell(row=2, column=col_idx, value=header)
    for fila_idx, fila in enumerate(filas, start=3):
        for col_idx, valor in enumerate(fila, start=1):
            ws.cell(row=fila_idx, column=col_idx, value=valor)
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.read()


def _parsear(headers: List[str], filas: List[List[Any]]) -> Dict[str, Any]:
    excel_bytes = _crear_excel_detalle(headers, filas)
    excel_file = pd.ExcelFile(io.BytesIO(excel_bytes))
    return PedidoParserDetalle().parsear_nuevo_formato_v1_0(excel_file, excel_bytes, "detalle")


def _info_producto_de(resultado: Dict[str, Any], codigo_producto: str) -> Optional[Dict[str, Any]]:
    for item in resultado["data"]:
        if item["codigoProducto"] == codigo_producto and item.get("info_producto"):
            return item["info_producto"]
    return None


# Headers base: producto (A-I) + bodegas (J-L) + tiendas (M+), con "Costo" al final.
_HEADERS_BASE = [
    "Codigo Producto", "Codigo de Barras", "Descripcion SAP", "Pack",
    "Costo sin IVA", "Descuento", "PVP", "Subtotal", "Total",
    "RL-CD", "OUT-CD", "RG-CD",
    "ABC01", "ABC02",
    "Costo",
]


class TestColumnaCostoMargen:
    def test_costo_sin_iva_sigue_mapeando_a_costo_doc_sin_iva(self):
        """RN-05: la columna preexistente "Costo sin IVA" no cambia de comportamiento."""
        fila = ["SKU-1", "7501234567890", "Producto Uno", "6", 10.5, 0, 15.0, "", "", "", "", "", 10, "", 3.2]
        resultado = _parsear(_HEADERS_BASE, [fila])

        info = _info_producto_de(resultado, "SKU-1")
        assert info is not None
        assert info["costo_doc_sin_iva"] == 10.5

    def test_columna_costo_mapea_a_costo_margen_distinto_de_costo_doc_sin_iva(self):
        """La nueva columna "Costo" es un campo aparte (costo_margen), no pisa costo_doc_sin_iva."""
        fila = ["SKU-1", "7501234567890", "Producto Uno", "6", 10.5, 0, 15.0, "", "", "", "", "", 10, "", 3.2]
        resultado = _parsear(_HEADERS_BASE, [fila])

        info = _info_producto_de(resultado, "SKU-1")
        assert info is not None
        assert info["costo_margen"] == 3.2
        assert info["costo_doc_sin_iva"] == 10.5
        assert info["costo_margen"] != info["costo_doc_sin_iva"]

    def test_columna_costo_ausente_no_rompe_el_parseo(self):
        """Excel sin columna "Costo" (plantillas anteriores a este requerimiento): sigue funcionando."""
        headers_sin_costo = _HEADERS_BASE[:-1]  # quita "Costo"
        fila = ["SKU-1", "7501234567890", "Producto Uno", "6", 10.5, 0, 15.0, "", "", "", "", "", 10, ""]
        resultado = _parsear(headers_sin_costo, [fila])

        info = _info_producto_de(resultado, "SKU-1")
        assert info is not None
        assert info.get("costo_margen") is None
        assert info["costo_doc_sin_iva"] == 10.5

    def test_columna_costo_al_final_no_rompe_deteccion_de_bodegas_ni_tiendas(self):
        """La columna "Costo" al final de la hoja no debe interferir con bodegas/tiendas."""
        fila = ["SKU-1", "7501234567890", "Producto Uno", "6", 10.5, 0, 15.0, "", "", 4, 2, 1, 10, 5, 3.2]
        resultado = _parsear(_HEADERS_BASE, [fila])

        distribuciones = {d["idTienda"]: d["cantidad"] for d in resultado["data"] if d["codigoProducto"] == "SKU-1"}
        # 3 bodegas + 2 tiendas = 5 distribuciones; ninguna debe perderse por la columna Costo.
        assert distribuciones.get("RL-CD") == 4
        assert distribuciones.get("OUT-CD") == 2
        assert distribuciones.get("RG-CD") == 1

    def test_columna_costo_unitario_en_medio_no_pierde_bodega_regions(self):
        """
        Plantilla real (incidente 2026-08-07): "Costo Unitario" se inserta en la columna H
        (después de PVP, antes de Total/Subtotal), corriendo bodegas de J-L a K-M. Antes del
        fix, RG-CD quedaba fuera del rango fijo INDICES_BODEGAS=[9,10,11] y se perdía en
        silencio (ni bodega ni tienda). Reproduce exactamente ese layout.
        """
        headers = [
            "Codigo Producto", "Codigo de Barras", "Descripcion SAP", "Pack",
            "Costo sin IVA", "Descuento", "PVP",
            "Costo Unitario",
            "Total", "Subtotal",
            "RL-CD", "OUT-CD", "RG-CD",
            "ABC01", "ABC02",
        ]
        fila = ["SKU-1", "7501234567890", "Producto Uno", "6", 10.5, 0, 15.0, 3.2, "", "", 4, 2, 1, 10, 5]
        resultado = _parsear(headers, [fila])

        info = _info_producto_de(resultado, "SKU-1")
        assert info is not None
        assert info["costo_margen"] == 3.2
        assert info["costo_doc_sin_iva"] == 10.5
        # La misma columna física "Costo Unitario" también se lee aparte como
        # costo_unitario_excel, para contrastarla contra el cálculo del backend
        # (ver PedidoService._resolver_costo_unitario_temprana) y detectar fórmulas
        # mal aplicadas en esa celda.
        assert info["costo_unitario_excel"] == 3.2

        distribuciones = {d["idTienda"]: d["cantidad"] for d in resultado["data"] if d["codigoProducto"] == "SKU-1"}
        assert distribuciones.get("RL-CD") == 4
        assert distribuciones.get("OUT-CD") == 2
        assert distribuciones.get("RG-CD") == 1
        assert distribuciones.get("ABC01") == 10
        assert distribuciones.get("ABC02") == 5
        assert distribuciones.get("ABC01") == 10
        assert distribuciones.get("ABC02") == 5

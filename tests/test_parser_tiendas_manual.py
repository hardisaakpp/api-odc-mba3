"""
Validación del método leer_hoja_tiendas() - Test manual
Prueba la normalización de datos desde Excel
"""

import io
import openpyxl
from src.database.services.pedido_parser import PedidoParser

def crear_excel_test_tiendas():
    """Crea un Excel de test con hoja Tiendas."""
    from openpyxl import Workbook
    
    wb = Workbook()
    ws = wb.active
    assert ws is not None, "No se pudo crear la hoja activa"
    ws.title = "Tiendas"
    
    # Headers
    ws['A1'] = 'Código'
    ws['B1'] = 'Nombre'
    ws['C1'] = 'Activa'
    
    # Datos (con variaciones: espacios, minúsculas, diferentes booleanos)
    ws['A2'] = '  TRL-QSN  '  # Con espacios
    ws['B2'] = 'ROLAND QUICENTRO NORTE'
    ws['C2'] = 'SI'
    
    ws['A3'] = 'trl-cci'  # Minúsculas
    ws['B3'] = 'ROLAND CCI'
    ws['C3'] = 'yes'  # Booleano en inglés
    
    ws['A4'] = 'TRL-OLD'
    ws['B4'] = 'ROLAND ANTIGUA'
    ws['C4'] = 'NO'  # Inactiva
    
    ws['A5'] = 'TRL-TEST'
    ws['B5'] = 'TIENDA TEST'
    ws['C5'] = 'false'  # Booleano False
    
    # Guardar en bytes
    excel_bytes = io.BytesIO()
    wb.save(excel_bytes)
    excel_bytes.seek(0)
    return excel_bytes.getvalue()


def test_parser_leer_hoja_tiendas():
    """Test manual del parser."""
    parser = PedidoParser()
    excel_bytes = crear_excel_test_tiendas()
    
    # Act
    catalogo = parser.leer_hoja_tiendas(excel_bytes)
    
    # Assert
    print("\n📊 Catálogo leído:")
    for codigo, info in catalogo.items():
        print(f"  {codigo}: {info['nombre']} (Activa: {info['activa']})")
    
    # Validaciones
    assert len(catalogo) == 4, f"Esperaba 4 tiendas, obtuve {len(catalogo)}"
    
    # Verificar normalización de espacios y mayúsculas
    assert "TRL-QSN" in catalogo, "TRL-QSN no fue normalizado correctamente"
    assert catalogo["TRL-QSN"]["activa"] is True, "TRL-QSN debe estar activa"
    
    # Verificar minúsculas
    assert "TRL-CCI" in catalogo, "TRL-CCI no fue normalizado correctamente"
    assert catalogo["TRL-CCI"]["activa"] is True, "TRL-CCI debe estar activa"
    
    # Verificar inactivas
    assert "TRL-OLD" in catalogo, "TRL-OLD no se cargó"
    assert catalogo["TRL-OLD"]["activa"] is False, "TRL-OLD debe estar inactiva"
    
    assert "TRL-TEST" in catalogo, "TRL-TEST no se cargó"
    assert catalogo["TRL-TEST"]["activa"] is False, "TRL-TEST debe estar inactiva"
    
    print("✅ Test manual PASÓ: Normalización funciona correctamente")
    print(f"   - 4 tiendas cargadas")
    print(f"   - 2 activas: TRL-QSN, TRL-CCI")
    print(f"   - 2 inactivas: TRL-OLD, TRL-TEST")


if __name__ == "__main__":
    test_parser_leer_hoja_tiendas()

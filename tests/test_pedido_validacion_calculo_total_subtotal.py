"""
Tests para la validación bloqueante de TOTAL/SUBTOTAL en la carga de plantilla Excel.

Incidente 2026-08-28 (pedido PDES-08.42.2026.CJ): una fórmula de SUBTOTAL con la fila del
INDICE() desplazada en +1 (=SUMA(N3:INDICE(4:4,...)) en vez de INDICE(3:3,...)) hacía que
cada fila sumara también las cantidades de tienda de la fila siguiente. El parser ya detecta
esto comparando el valor leído de la celda contra la suma real de bodegas/tiendas (independiente
de la fórmula), pero antes solo generaba un WARNING no bloqueante que nunca llegaba al usuario
(se descartaba en pedido_service.py antes de armar la respuesta HTTP).

Cubre:
- El parser marca la fila como ERROR (no WARNING) cuando SUBTOTAL o TOTAL no coinciden con las
  cantidades reales de tienda/bodega.
- PedidoService._validar_calculo_total_subtotal lanza BusinessException bloqueante cuando hay
  errores de cálculo, y no lanza nada cuando las fórmulas son correctas.
"""

import io
from typing import Any, Dict, List

import pandas as pd
import pytest
from openpyxl import Workbook

from src.database.services.pedido_parser.pedido_parser_detalle import PedidoParserDetalle
from src.database.services.pedido_service import PedidoService
from src.exceptions import BusinessException

# Codigo Producto, Codigo Barras, Descripcion SAP, Pack, Costo sin IVA, Descuento, PVP,
# Total, Subtotal, RL-CD, OUT-CD, RG-CD, luego columnas de tienda (ABC01/ABC02).
_HEADERS = [
    "Codigo Producto", "Codigo de Barras", "Descripcion SAP", "Pack",
    "Costo sin IVA", "Descuento", "PVP", "Total", "Subtotal",
    "RL-CD", "OUT-CD", "RG-CD",
    "ABC01", "ABC02",
]


def _crear_excel_detalle(filas: List[List[Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "detalle"
    for col_idx, header in enumerate(_HEADERS, start=1):
        ws.cell(row=1, column=col_idx, value=header)
        ws.cell(row=2, column=col_idx, value=header)
    for fila_idx, fila in enumerate(filas, start=3):
        for col_idx, valor in enumerate(fila, start=1):
            ws.cell(row=fila_idx, column=col_idx, value=valor)
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.read()


def _parsear(filas: List[List[Any]]) -> Dict[str, Any]:
    excel_bytes = _crear_excel_detalle(filas)
    excel_file = pd.ExcelFile(io.BytesIO(excel_bytes))
    return PedidoParserDetalle().parsear_nuevo_formato_v1_0(excel_file, excel_bytes, "detalle")


def _errores_de_campo(resultado: Dict[str, Any], campo: str) -> List[Dict[str, Any]]:
    return [e for e in resultado["errores"] if isinstance(e, dict) and e.get("campo") == campo]


class TestParserDeteccionCalculoTotalSubtotal:
    def test_total_y_subtotal_correctos_no_generan_error(self):
        # bodegas: 4+2+1=7, tiendas: 10+5=15 -> subtotal=15, total=15+7=22
        fila = ["SKU-1", "7501234567890", "Producto Uno", "6", 10.5, 0, 15.0, 22, 15, 4, 2, 1, 10, 5]
        resultado = _parsear([fila])

        assert _errores_de_campo(resultado, "TOTAL") == []
        assert _errores_de_campo(resultado, "SUBTOTAL") == []

    def test_subtotal_desfasado_genera_error_bloqueante_tipo_error(self):
        # Reproduce el bug real: SUBTOTAL trae 20 (15 reales + 5 "prestados" de la fila siguiente).
        fila = ["SKU-1", "7501234567890", "Producto Uno", "6", 10.5, 0, 15.0, 22, 20, 4, 2, 1, 10, 5]
        resultado = _parsear([fila])

        errores_subtotal = _errores_de_campo(resultado, "SUBTOTAL")
        assert len(errores_subtotal) == 1
        assert errores_subtotal[0]["tipo"] == "ERROR"
        assert errores_subtotal[0]["fila"] == 4
        assert errores_subtotal[0]["codigoProducto"] == "SKU-1"
        assert "no coincide" in errores_subtotal[0]["mensaje"]
        # TOTAL sigue correcto en este caso (22 == 15 tiendas + 7 bodegas)
        assert _errores_de_campo(resultado, "TOTAL") == []

    def test_total_incorrecto_genera_error_bloqueante_tipo_error(self):
        fila = ["SKU-1", "7501234567890", "Producto Uno", "6", 10.5, 0, 15.0, 30, 15, 4, 2, 1, 10, 5]
        resultado = _parsear([fila])

        errores_total = _errores_de_campo(resultado, "TOTAL")
        assert len(errores_total) == 1
        assert errores_total[0]["tipo"] == "ERROR"
        assert errores_total[0]["fila"] == 4
        assert errores_total[0]["codigoProducto"] == "SKU-1"


class TestServicioValidacionCalculoBloqueaCarga:
    def test_sin_errores_de_calculo_no_lanza_excepcion(self):
        fila = ["SKU-1", "7501234567890", "Producto Uno", "6", 10.5, 0, 15.0, 22, 15, 4, 2, 1, 10, 5]
        resultado = _parsear([fila])

        PedidoService._validar_calculo_total_subtotal(resultado)  # no debe lanzar

    def test_con_errores_de_calculo_lanza_business_exception_con_detalle(self):
        fila = ["SKU-1", "7501234567890", "Producto Uno", "6", 10.5, 0, 15.0, 22, 20, 4, 2, 1, 10, 5]
        resultado = _parsear([fila])

        with pytest.raises(BusinessException) as exc_info:
            PedidoService._validar_calculo_total_subtotal(resultado)

        exc = exc_info.value
        assert exc.error_code == "CALCULO_TOTAL_SUBTOTAL_INCORRECTO"
        assert "SKU-1" in exc.message
        assert "Fila 4" in exc.message
        calculo_campos = exc.data["validacion"]["calculo_campos"]
        assert calculo_campos["total_filas_con_error"] == 1
        assert calculo_campos["errores"][0]["campo"] == "SUBTOTAL"

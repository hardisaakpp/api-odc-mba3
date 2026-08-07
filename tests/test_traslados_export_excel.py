"""Tests del export a Excel de Solicitudes de Traslado (requerimiento 1.9)."""

from typing import List, Optional

from fastapi import FastAPI
from fastapi.routing import APIRoute
from openpyxl import load_workbook

from src.endpoints.pedidos.pedidos_redistribucion import register_redistribucion_endpoints
from src.endpoints.pedidos.redistribucion._excel_traslados import (
    SIN_TRASLADO_BODEGA,
    SIN_TRASLADO_PENDIENTE,
    _FILA_DATOS_INICIO,
    _FILA_DOCNUM,
    _FILA_SUBHEADER,
    _FILA_TIENDA,
    agrupar_filas_por_producto,
    construir_matriz_traslados,
    generar_excel_traslados,
)
from src.security.jwt_auth import verify_jwt_token

RUTA_EXPORT = "/api/pedidos/{pedido_id}/traslados/exportar-excel"


def _producto(
    codigo_interno: str,
    codigo_barras: str,
    distribuciones: List[dict],
    confirmado_total: float = 0.0,
) -> dict:
    return {
        "codigo_interno": codigo_interno,
        "descripcion_sap": f"DESCRIPCION {codigo_interno}",
        "codigo_barras": codigo_barras,
        "cantidad_confirmada_total": confirmado_total,
        "distribuciones": distribuciones,
    }


def _dist(
    tienda_codigo: str,
    tienda_nombre: str,
    solicitada: float,
    confirmada: float,
    asignada: float,
    es_bodega: bool = False,
) -> dict:
    return {
        "tienda_codigo": tienda_codigo,
        "tienda_nombre": tienda_nombre,
        "cantidad_solicitada": solicitada,
        "cantidad_confirmada": confirmada,
        "cantidad_asignada": asignada,
        "es_bodega": es_bodega,
    }


class TestConstruirMatrizTraslados:
    """La función pura que arma los datos del export."""

    def test_una_fila_por_producto_y_tienda(self):
        productos = [
            _producto("5986.NG.6", "7861224907024", [
                _dist("TRL-CON", "ROLAND CONDADO", 7, 7, 7),
                _dist("TRL-SMA", "ROLAND SAN LUIS", 8, 8, 8),
            ]),
            _producto("8234.2", "7861105846473", [
                _dist("TRL-CON", "ROLAND CONDADO", 7, 5, 5),
            ]),
        ]
        matriz = construir_matriz_traslados(productos, {})

        assert len(matriz["filas"]) == 3
        assert len(matriz["tiendas"]) == 2

    def test_doc_num_se_resuelve_normalizando_el_codigo(self):
        """El mapa viene normalizado (UPPER+TRIM); el lookup debe normalizar igual."""
        productos = [
            _producto("5986.NG.6", "7861224907024", [
                _dist(" trl-con ", "ROLAND CONDADO", 7, 7, 7),
            ]),
        ]
        matriz = construir_matriz_traslados(productos, {"TRL-CON": 70236990})

        assert matriz["filas"][0]["doc_num"] == 70236990

    def test_tienda_sin_itr_queda_registrada_como_pendiente(self):
        productos = [
            _producto("5986.NG.6", "7861224907024", [
                _dist("TRL-CON", "ROLAND CONDADO", 7, 7, 7),
                _dist("TRL-SMA", "ROLAND SAN LUIS", 8, 8, 8),
            ]),
        ]
        matriz = construir_matriz_traslados(productos, {"TRL-CON": 70236990})

        assert matriz["tiendas_sin_doc_num"] == ["TRL-SMA"]

    def test_bodegas_van_al_final_y_no_cuentan_como_pendientes(self):
        """Las bodegas nunca tienen doc_num: no deben ensuciar la lista de pendientes."""
        productos = [
            _producto("5986.NG.6", "7861224907024", [
                _dist("RL-CD", "CEDI ROLAND", 100, 100, 100, es_bodega=True),
                _dist("TRL-CON", "ROLAND CONDADO", 7, 7, 7),
            ]),
        ]
        matriz = construir_matriz_traslados(productos, {"TRL-CON": 70236990})

        assert [t["codigo"] for t in matriz["tiendas"]] == ["TRL-CON", "RL-CD"]
        assert matriz["tiendas_sin_doc_num"] == []
        assert matriz["bodegas"] == ["RL-CD"]

    def test_totales_por_tienda_y_gran_total(self):
        productos = [
            _producto("5986.NG.6", "7861224907024", [
                _dist("TRL-CON", "ROLAND CONDADO", 7, 7, 7),
            ]),
            _producto("8234.2", "7861105846473", [
                _dist("TRL-CON", "ROLAND CONDADO", 8, 6, 5),
            ]),
        ]
        matriz = construir_matriz_traslados(productos, {})

        assert matriz["totales_por_tienda"]["TRL-CON"] == {
            "solicitada": 15.0,
            "confirmada": 13.0,
            "asignada": 12.0,
        }
        assert matriz["gran_total_solicitada"] == 15.0
        assert matriz["gran_total_confirmada"] == 13.0
        assert matriz["gran_total_asignada"] == 12.0

    def test_cantidades_nulas_se_tratan_como_cero(self):
        productos = [
            _producto("5986.NG.6", "7861224907024", [
                {
                    "tienda_codigo": "TRL-CON",
                    "tienda_nombre": "ROLAND CONDADO",
                    "cantidad_solicitada": None,
                    "cantidad_confirmada": None,
                    "cantidad_asignada": None,
                },
            ]),
        ]
        matriz = construir_matriz_traslados(productos, {})

        assert matriz["gran_total_solicitada"] == 0.0

    def test_distribucion_sin_codigo_de_tienda_se_ignora(self):
        productos = [
            _producto("5986.NG.6", "7861224907024", [
                _dist("", "SIN CODIGO", 7, 7, 7),
                _dist("TRL-CON", "ROLAND CONDADO", 7, 7, 7),
            ]),
        ]
        matriz = construir_matriz_traslados(productos, {})

        assert len(matriz["filas"]) == 1
        assert matriz["filas"][0]["tienda_codigo"] == "TRL-CON"

    def test_producto_sin_distribuciones_no_rompe(self):
        matriz = construir_matriz_traslados([_producto("X", "1", [])], {})

        assert matriz["filas"] == []
        assert matriz["tiendas"] == []

    def test_confirmado_total_del_producto_se_propaga_a_sus_filas(self):
        """El confirmado real del proveedor solo existe a nivel producto."""
        productos = [
            _producto("5986.NG.6", "7861224907024", [
                _dist("TRL-CON", "ROLAND CONDADO", 7, 7, 7),
                _dist("TRL-SMA", "ROLAND SAN LUIS", 8, 8, 8),
            ], confirmado_total=15.0),
        ]
        matriz = construir_matriz_traslados(productos, {})

        assert all(f["confirmado_total_producto"] == 15.0 for f in matriz["filas"])


class TestGenerarExcelTraslados:
    """El libro Excel resultante."""

    @staticmethod
    def _libro(productos, itr_por_tienda=None):
        buffer = generar_excel_traslados(
            pedido_id="11111111-2222-3333-4444-555555555555",
            pedido_data={"cabecera": {"numero_pedido": "PROL-06.76.2026"}, "detalles": []},
            distribucion_data={"version_actual": 2},
            productos_enriquecidos=productos,
            user={"username": "cortiz"},
            itr_por_tienda=itr_por_tienda or {},
        )
        return load_workbook(buffer)

    def test_hojas_esperadas(self):
        wb = self._libro([
            _producto("5986.NG.6", "7861224907024", [
                _dist("TRL-CON", "ROLAND CONDADO", 7, 7, 7),
            ]),
        ])

        assert wb.sheetnames == [
            "Información del Pedido",
            "Matriz por Tienda",
            "Traslados (Detalle)",
        ]

    def test_encabezados_del_requerimiento_presentes(self):
        """Las 5 columnas que el requerimiento pide como mínimo."""
        wb = self._libro([
            _producto("5986.NG.6", "7861224907024", [
                _dist("TRL-CON", "ROLAND CONDADO", 7, 7, 7),
            ]),
        ])
        ws = wb["Traslados (Detalle)"]
        headers = [ws.cell(row=3, column=c).value for c in range(1, 11)]

        for requerido in (
            "Número de Traslado",
            "Nombre de la Tienda",
            "Código de Barras",
            "Cantidad Solicitada",
            "Cantidad Confirmada",
        ):
            assert requerido in headers

    def test_fila_de_datos_con_doc_num(self):
        wb = self._libro(
            [
                _producto("5986.NG.6", "7861224907024", [
                    _dist("TRL-CON", "ROLAND CONDADO", 7, 6, 5),
                ], confirmado_total=6.0),
            ],
            {"TRL-CON": 70236990},
        )
        ws = wb["Traslados (Detalle)"]

        assert ws.cell(row=4, column=1).value == 70236990
        assert ws.cell(row=4, column=2).value == "TRL-CON"
        assert ws.cell(row=4, column=3).value == "ROLAND CONDADO"
        assert ws.cell(row=4, column=6).value == "7861224907024"
        assert ws.cell(row=4, column=7).value == 7
        assert ws.cell(row=4, column=8).value == 6
        assert ws.cell(row=4, column=9).value == 5

    def test_bodega_muestra_no_aplica_en_numero_de_traslado(self):
        wb = self._libro([
            _producto("5986.NG.6", "7861224907024", [
                _dist("RL-CD", "CEDI ROLAND", 100, 100, 100, es_bodega=True),
            ]),
        ])
        ws = wb["Traslados (Detalle)"]

        assert ws.cell(row=4, column=1).value == SIN_TRASLADO_BODEGA

    def test_tienda_sin_itr_muestra_pendiente(self):
        wb = self._libro([
            _producto("5986.NG.6", "7861224907024", [
                _dist("TRL-CON", "ROLAND CONDADO", 7, 7, 7),
            ]),
        ])
        ws = wb["Traslados (Detalle)"]

        assert ws.cell(row=4, column=1).value == SIN_TRASLADO_PENDIENTE

    def test_encabezado_congelado_para_desplazamiento(self):
        wb = self._libro([
            _producto("5986.NG.6", "7861224907024", [
                _dist("TRL-CON", "ROLAND CONDADO", 7, 7, 7),
            ]),
        ])

        assert wb["Traslados (Detalle)"].freeze_panes == "A4"

    def test_pedido_sin_distribuciones_genera_libro_valido(self):
        """No debe reventar: el botón no se gatea por estado."""
        wb = self._libro([_producto("A", "1", [])])
        ws = wb["Traslados (Detalle)"]

        assert ws.cell(row=3, column=1).value == "Número de Traslado"

    def test_cabecera_declara_el_origen_del_numero_de_traslado(self):
        """El archivo debe explicitar de dónde sale el dato y sus límites."""
        wb = self._libro([
            _producto("A", "1", [_dist("TRL-CON", "ROLAND CONDADO", 7, 7, 7)]),
        ])
        ws = wb["Información del Pedido"]
        etiquetas = [ws.cell(row=r, column=1).value for r in range(3, ws.max_row + 1)]

        assert "Origen del Nº de traslado" in etiquetas
        assert "Alcance del Nº de traslado" in etiquetas


class TestAgruparFilasPorProducto:
    """Pivotea las filas planas (producto, tienda) a una fila por producto.

    Ni lo solicitado ni lo confirmado se conservan aquí: las celdas por tienda
    (y su suma) reflejan la redistribución (cantidad_asignada), la misma que
    se ve por tienda en la pantalla de "Ajuste y Redistribución a Tiendas".
    """

    def test_una_fila_por_producto_con_celdas_por_tienda(self):
        matriz = construir_matriz_traslados(
            [
                _producto("5986.NG.6", "7861224907024", [
                    _dist("TRL-CON", "ROLAND CONDADO", 7, 7, 6),
                    _dist("TRL-SMA", "ROLAND SAN LUIS", 8, 6, 9),
                ]),
            ],
            {},
        )
        productos = agrupar_filas_por_producto(matriz["filas"])

        assert len(productos) == 1
        celdas = productos[0]["celdas"]
        assert celdas["TRL-CON"] == 6
        assert celdas["TRL-SMA"] == 9

    def test_totales_por_producto(self):
        matriz = construir_matriz_traslados(
            [
                _producto("5986.NG.6", "7861224907024", [
                    _dist("TRL-CON", "ROLAND CONDADO", 7, 7, 6),
                    _dist("TRL-SMA", "ROLAND SAN LUIS", 8, 6, 9),
                ], confirmado_total=13.0),
            ],
            {},
        )
        productos = agrupar_filas_por_producto(matriz["filas"])

        assert productos[0]["total_asignada"] == 15

    def test_conserva_orden_de_primera_aparicion(self):
        """agrupar_filas_por_producto no reordena: respeta el orden de la lista recibida."""
        fila_base = {
            "doc_num": None, "tienda_codigo": "TRL-CON", "tienda_nombre": "ROLAND CONDADO",
            "es_bodega": False, "descripcion": "", "cantidad_solicitada": 1,
            "cantidad_confirmada": 1, "cantidad_asignada": 1, "confirmado_total_producto": 1,
        }
        filas = [
            {**fila_base, "codigo_articulo": "B", "codigo_barras": "2"},
            {**fila_base, "codigo_articulo": "A", "codigo_barras": "1"},
        ]

        productos = agrupar_filas_por_producto(filas)

        assert [p["codigo_articulo"] for p in productos] == ["B", "A"]


class TestHojaMatrizTraslados:
    """La hoja 'Matriz por Tienda': tiendas en columnas con cantidad_asignada
    (la redistribución), 'Total General' como una sola columna con la suma."""

    @staticmethod
    def _libro(productos, itr_por_tienda=None, numero_pedido="PROL-06.76.2026"):
        buffer = generar_excel_traslados(
            pedido_id="11111111-2222-3333-4444-555555555555",
            pedido_data={"cabecera": {"numero_pedido": numero_pedido}, "detalles": []},
            distribucion_data={"version_actual": 1},
            productos_enriquecidos=productos,
            user={"username": "cortiz"},
            itr_por_tienda=itr_por_tienda or {},
        )
        return load_workbook(buffer)["Matriz por Tienda"]

    def test_numero_de_pedido_en_a1(self):
        ws = self._libro(
            [_producto("A", "1", [_dist("TRL-CON", "ROLAND CONDADO", 1, 1, 1)])],
            numero_pedido="PROL-06.76.2026",
        )

        assert ws.cell(row=1, column=1).value == "PROL-06.76.2026"

    def test_docnum_como_encabezado_de_bloque_de_tienda(self):
        # 1 tienda (columna 3: DocNum + nombre fusionado con la fila de
        # subencabezado) + Total General (col 4, sin subencabezado)
        ws = self._libro(
            [_producto("A", "1", [_dist("TRL-CON", "ROLAND CONDADO", 1, 1, 1)])],
            itr_por_tienda={"TRL-CON": 70236990},
        )

        assert ws.cell(row=_FILA_DOCNUM, column=3).value == 70236990
        assert ws.cell(row=_FILA_TIENDA, column=3).value == "ROLAND CONDADO"
        assert ws.cell(row=_FILA_SUBHEADER, column=3).value is None
        assert ws.cell(row=_FILA_DOCNUM, column=4).value == "Total General"
        assert ws.cell(row=_FILA_SUBHEADER, column=4).value is None

    def test_bodega_sin_docnum_muestra_no_aplica(self):
        ws = self._libro(
            [_producto("A", "1", [_dist("RL-CD", "CEDI ROLAND", 5, 5, 5, es_bodega=True)])],
        )

        assert ws.cell(row=_FILA_DOCNUM, column=3).value == SIN_TRASLADO_BODEGA

    def test_tienda_sin_itr_muestra_pendiente(self):
        ws = self._libro(
            [_producto("A", "1", [_dist("TRL-CON", "ROLAND CONDADO", 1, 1, 1)])],
        )

        assert ws.cell(row=_FILA_DOCNUM, column=3).value == SIN_TRASLADO_PENDIENTE

    def test_por_tienda_muestra_cantidad_asignada(self):
        """La celda por tienda es la redistribución (asignada), no lo solicitado ni lo confirmado."""
        ws = self._libro(
            [_producto("A", "1", [_dist("TRL-CON", "ROLAND CONDADO", 7, 6, 5)], confirmado_total=6.0)],
        )

        assert ws.cell(row=_FILA_DATOS_INICIO, column=3).value == 5  # Asignada tienda
        assert ws.cell(row=_FILA_DATOS_INICIO, column=4).value == 5  # Total General (suma de asignada)

    def test_columna_total_general_es_una_sola_y_suma_la_redistribucion(self):
        ws = self._libro(
            [
                _producto("A", "1", [
                    _dist("TRL-CON", "ROLAND CONDADO", 7, 7, 6),
                    _dist("TRL-SMA", "ROLAND SAN LUIS", 8, 6, 9),
                ], confirmado_total=13.0),
            ],
        )
        # 2 tiendas x 1 columna + 2 fijas = col 5 (Total General, única columna)
        assert ws.cell(row=_FILA_DATOS_INICIO, column=3).value == 6
        assert ws.cell(row=_FILA_DATOS_INICIO, column=4).value == 9
        assert ws.cell(row=_FILA_DATOS_INICIO, column=5).value == 15

    def test_fila_totales_acumula_por_tienda_y_general(self):
        ws = self._libro(
            [
                _producto("A", "1", [_dist("TRL-CON", "ROLAND CONDADO", 7, 7, 6)], confirmado_total=7.0),
                _producto("B", "2", [_dist("TRL-CON", "ROLAND CONDADO", 8, 6, 4)], confirmado_total=6.0),
            ],
        )
        fila_totales = _FILA_DATOS_INICIO + 2  # 2 productos, luego TOTALES

        assert ws.cell(row=fila_totales, column=1).value == "TOTALES"
        assert ws.cell(row=fila_totales, column=3).value == 10  # Asignada TRL-CON (única columna de la tienda)
        assert ws.cell(row=fila_totales, column=4).value == 10  # Total General (suma de asignada)

    def test_encabezado_y_columnas_fijas_congelados(self):
        ws = self._libro(
            [_producto("A", "1", [_dist("TRL-CON", "ROLAND CONDADO", 1, 1, 1)])],
        )

        assert ws.freeze_panes == "C6"

    def test_multiples_tiendas_no_pisan_columnas(self):
        """Regresión: cada tienda debe ocupar su propia columna (asignada), sin solaparse."""
        ws = self._libro(
            [
                _producto("A", "1", [
                    _dist("TRL-CON", "ROLAND CONDADO", 1, 1, 1),
                    _dist("TRL-QSN", "ROLAND QUICENTRO NORTE", 2, 2, 2),
                    _dist("TRL-SMA", "ROLAND SAN LUIS", 3, 3, 3),
                ]),
            ],
        )

        valores = [ws.cell(row=_FILA_DATOS_INICIO, column=c).value for c in range(3, 6)]
        assert valores == [1, 2, 3]


class TestContratoRutaExportarTraslados:
    """Contrato HTTP de la ruta nueva."""

    @staticmethod
    def _routes() -> List[APIRoute]:
        app = FastAPI()
        register_redistribucion_endpoints(app)
        return [r for r in app.routes if isinstance(r, APIRoute)]

    def _ruta(self) -> Optional[APIRoute]:
        for route in self._routes():
            if route.path == RUTA_EXPORT and "GET" in route.methods:
                return route
        return None

    def test_ruta_registrada(self):
        assert self._ruta() is not None

    def test_recibe_pedido_id_y_version(self):
        route = self._ruta()

        assert {p.name for p in route.dependant.path_params} == {"pedido_id"}
        assert "version" in {p.name for p in route.dependant.query_params}

    def test_exige_autenticacion(self):
        route = self._ruta()
        dependencias = {d.call for d in route.dependant.dependencies}

        assert verify_jwt_token in dependencias

    def test_autoriza_a_ambos_roles(self):
        """El requerimiento pide que MABEL y EL RAYO puedan descargar."""
        route = self._ruta()
        cierres = [
            d.call.__closure__
            for d in route.dependant.dependencies
            if d.call is not verify_jwt_token and d.call.__closure__
        ]
        permisos = {
            perm
            for closure in cierres
            for celda in closure
            if isinstance(celda.cell_contents, list)
            for perm in celda.cell_contents
        }

        assert {"mabel.view", "odcsap.view"} <= permisos

    def test_no_altera_la_ruta_de_descarga_existente(self):
        """Guarda de no regresión: el export de distribución sigue registrado."""
        paths = {r.path for r in self._routes()}

        assert "/api/pedidos/{pedido_id}/distribucion/descargar-excel" in paths

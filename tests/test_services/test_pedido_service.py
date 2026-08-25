"""
Tests para PedidoService - Parsing de Excel (formato v1.0).

El formato v1.0 espera:
- Hoja "cabecera": formato vertical Campo | Valor (o horizontal)
- Hoja "detalle": formato ancho con columnas de producto (A-I), bodegas (J-L), tiendas (M+)
  - Fila 1: Nombres de columnas (header)
  - Fila 2: Códigos de columnas (para tiendas/bodegas)
  - Fila 3+: Datos de productos
"""

import io
from typing import Optional
from openpyxl import Workbook

from src.database.services.pedido_service import PedidoService


# Columnas de producto estándar (A-I, índices 0-8)
COLUMNAS_PRODUCTO = [
    "Codigo Producto",       # A (0)
    "Codigo de Barras",      # B (1)
    "Descripcion SAP",       # C (2)
    "Pack",                  # D (3)
    "Costo sin IVA",         # E (4)
    "Descuento",             # F (5)
    "PVP",                   # G (6)
    "Subtotal",              # H (7)
    "Total",                 # I (8)
]

# Columnas de bodegas (J-L, índices 9-11)
COLUMNAS_BODEGAS = [
    "RL-CD",    # J (9)
    "OUT-CD",   # K (10)
    "RG-CD",    # L (11)
]


def _crear_excel_v1(
    cabecera_campos: list,
    detalle_filas: list,
    tiendas: Optional[list[str]] = None,
    hoja_cabecera: str = "cabecera",
    hoja_detalle: str = "detalle",
    detalle_headers_override: Optional[list[str]] = None,
    omitir_hoja_detalle: bool = False,
) -> bytes:
    """
    Crea un Excel en formato v1.0 para testing.
    
    Args:
        cabecera_campos: Lista de tuplas (campo, valor) para la cabecera
        detalle_filas: Lista de dicts con datos por fila.
            Cada dict tiene: {"codigo": str, "tiendas": {"ABC01": 10, "ABC02": 5}}
            Opcionalmente: {"descripcion": str, "costo": float, ...}
        tiendas: Lista de códigos de tienda (ej: ["ABC01", "ABC02"])
        hoja_cabecera: Nombre de la hoja de cabecera
        hoja_detalle: Nombre de la hoja de detalle
        detalle_headers_override: Si se provee, usa estos headers en vez de los estándar
        omitir_hoja_detalle: Si True, no crea la hoja de detalle
    """
    if tiendas is None:
        tiendas = ["ABC01", "ABC02"]
    
    wb = Workbook()
    
    # === Hoja Cabecera ===
    ws_cab = wb.active
    if ws_cab is None:
        raise RuntimeError("No se pudo crear hoja de cabecera")
    ws_cab.title = hoja_cabecera
    
    # Header row
    ws_cab.cell(row=1, column=1, value="Campo")
    ws_cab.cell(row=1, column=2, value="Valor")
    ws_cab.cell(row=1, column=3, value="Ejemplo")
    
    for idx, (campo, valor) in enumerate(cabecera_campos, start=2):
        ws_cab.cell(row=idx, column=1, value=campo)
        ws_cab.cell(row=idx, column=2, value=valor)
        ws_cab.cell(row=idx, column=3, value=valor)
    
    # === Hoja Detalle ===
    if not omitir_hoja_detalle:
        ws_det = wb.create_sheet(title=hoja_detalle)
        
        if detalle_headers_override:
            # Headers personalizados (para test de columnas faltantes, case insensitive, etc.)
            for col_idx, header in enumerate(detalle_headers_override, start=1):
                ws_det.cell(row=1, column=col_idx, value=header)
            
            # Fila 2: códigos (mismos que headers para tiendas)
            for col_idx, header in enumerate(detalle_headers_override, start=1):
                ws_det.cell(row=2, column=col_idx, value=header)
            
            # Datos desde fila 3
            for fila_idx, fila_data in enumerate(detalle_filas, start=3):
                if isinstance(fila_data, list):
                    for col_idx, val in enumerate(fila_data, start=1):
                        ws_det.cell(row=fila_idx, column=col_idx, value=val)
                elif isinstance(fila_data, dict):
                    # Solo código en primera columna
                    ws_det.cell(row=fila_idx, column=1, value=fila_data.get("codigo"))
        else:
            # Formato estándar v1.0
            all_headers = COLUMNAS_PRODUCTO + COLUMNAS_BODEGAS + tiendas
            
            # Fila 1: Nombres de columnas
            for col_idx, header in enumerate(all_headers, start=1):
                ws_det.cell(row=1, column=col_idx, value=header)
            
            # Fila 2: Códigos de columnas (para bodegas y tiendas)
            for col_idx, header in enumerate(all_headers, start=1):
                if col_idx > len(COLUMNAS_PRODUCTO):  # Solo bodegas y tiendas
                    ws_det.cell(row=2, column=col_idx, value=header)
                else:
                    ws_det.cell(row=2, column=col_idx, value="")
            
            # Fila 3+: Datos de productos
            for fila_idx, fila_data in enumerate(detalle_filas, start=3):
                if fila_data is None:
                    continue  # Fila vacía
                    
                codigo = fila_data.get("codigo")
                ws_det.cell(row=fila_idx, column=1, value=codigo)
                
                # Campos opcionales de producto
                if "codigo_barras" in fila_data:
                    ws_det.cell(row=fila_idx, column=2, value=fila_data["codigo_barras"])
                if "descripcion" in fila_data:
                    ws_det.cell(row=fila_idx, column=3, value=fila_data["descripcion"])
                if "pack" in fila_data:
                    ws_det.cell(row=fila_idx, column=4, value=fila_data["pack"])
                if "costo" in fila_data:
                    ws_det.cell(row=fila_idx, column=5, value=fila_data["costo"])
                if "descuento" in fila_data:
                    ws_det.cell(row=fila_idx, column=6, value=fila_data["descuento"])
                if "pvp" in fila_data:
                    ws_det.cell(row=fila_idx, column=7, value=fila_data["pvp"])
                if "subtotal" in fila_data:
                    ws_det.cell(row=fila_idx, column=8, value=fila_data["subtotal"])
                if "total" in fila_data:
                    ws_det.cell(row=fila_idx, column=9, value=fila_data["total"])
                
                # Cantidades por tienda (columnas M+, índice 12+ = col 13+)
                tiendas_data = fila_data.get("tiendas", {})
                for tienda_idx, tienda_codigo in enumerate(tiendas):
                    col_tienda = len(COLUMNAS_PRODUCTO) + len(COLUMNAS_BODEGAS) + tienda_idx + 1
                    cantidad = tiendas_data.get(tienda_codigo)
                    if cantidad is not None:
                        ws_det.cell(row=fila_idx, column=col_tienda, value=cantidad)
    
    # Guardar a bytes
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.read()


class TestPedidoServiceParsearExcel:
    """Tests para parsear_excel_pedido() en formato v1.0."""
    
    def test_parsear_excel_formato_correcto(self):
        """Test que parsea un Excel con formato v1.0 correcto."""
        cabecera = [
            ("CODIGO PEDIDO", "P-12.48"),
            ("FECHA PEDIDO", "21/01/2025"),
            ("FECHA ENTREGA", "INMEDIATO"),
        ]
        
        detalle = [
            {"codigo": "R225.1012.BCO", "tiendas": {"ABC01": 10}},
            {"codigo": "A889.3321.NEG", "tiendas": {"ABC02": 5}},
        ]
        
        excel_bytes = _crear_excel_v1(cabecera, detalle, tiendas=["ABC01", "ABC02"])
        
        service = PedidoService()
        resultado = service.parsear_excel_pedido(excel_bytes)
        
        # Verificar estructura
        assert "cabecera" in resultado
        assert "resumen" in resultado
        assert "errores" in resultado
        assert "data" in resultado
        
        # Verificar cabecera: el parser mapea "CODIGO PEDIDO" → "numero_pedido"
        assert "numero_pedido" in resultado["cabecera"] or "CODIGO PEDIDO" in resultado["cabecera"]
        
        # Verificar resumen
        assert resultado["resumen"]["filasLeidas"] == 2
        assert resultado["resumen"]["filasValidas"] == 2
        assert resultado["resumen"]["filasInvalidas"] == 0
        
        # Verificar data (distribuciones por tienda)
        assert len(resultado["data"]) == 2
        
        # Verificar que se crearon distribuciones con datos correctos
        codigos = [d["codigoProducto"] for d in resultado["data"]]
        assert "R225.1012.BCO" in codigos
        assert "A889.3321.NEG" in codigos
        
        for item in resultado["data"]:
            assert item["cantidad"] > 0
            assert "filaOrigen" in item
    
    def test_parsear_excel_hoja_faltante(self):
        """Test que detecta cuando falta la hoja 'detalle'."""
        cabecera = [("CODIGO PEDIDO", "P-12.48")]
        
        excel_bytes = _crear_excel_v1(
            cabecera, [], tiendas=["ABC01"],
            omitir_hoja_detalle=True
        )
        
        service = PedidoService()
        resultado = service.parsear_excel_pedido(excel_bytes)
        
        # Debe tener error por hoja faltante
        assert len(resultado["errores"]) > 0
        assert any("detalle" in err["mensaje"].lower() for err in resultado["errores"])
        assert resultado["resumen"]["filasInvalidas"] > 0
    
    def test_parsear_excel_columnas_faltantes(self):
        """Test que detecta cuando falta la columna 'Codigo Producto' en detalle."""
        cabecera = [("CODIGO PEDIDO", "P-12.48")]
        
        # Detalle sin columna "Codigo Producto" (solo columnas irreconocibles)
        excel_bytes = _crear_excel_v1(
            cabecera, [{"codigo": "R225.1012.BCO"}], tiendas=["ABC01"],
            detalle_headers_override=["Columna Desconocida", "Otra Columna"]
        )
        
        service = PedidoService()
        resultado = service.parsear_excel_pedido(excel_bytes)
        
        # Debe tener errores por columna faltante
        assert len(resultado["errores"]) > 0
        assert any(
            "codigo producto" in err["mensaje"].lower() or
            "CODIGO PRODUCTO" in err["mensaje"]
            for err in resultado["errores"]
        )
    
    def test_parsear_excel_duplicados_detecta_ambos(self):
        """Test que procesa filas duplicadas (mismo producto, misma tienda) creando múltiples distribuciones."""
        cabecera = [("CODIGO PEDIDO", "P-12.48")]
        
        # Dos filas con mismo producto (en formato ancho, esto genera 2 distribuciones por tienda)
        detalle = [
            {"codigo": "R225.1012.BCO", "tiendas": {"ABC01": 10}},
            {"codigo": "R225.1012.BCO", "tiendas": {"ABC01": 5}},
        ]
        
        excel_bytes = _crear_excel_v1(cabecera, detalle, tiendas=["ABC01"])
        
        service = PedidoService()
        resultado = service.parsear_excel_pedido(excel_bytes)
        
        # Ambas filas se procesan como distribuciones separadas
        distribuciones_prod = [d for d in resultado["data"] if d["codigoProducto"] == "R225.1012.BCO"]
        assert len(distribuciones_prod) == 2
        cantidades = sorted([d["cantidad"] for d in distribuciones_prod])
        assert cantidades == [5, 10]
    
    def test_parsear_excel_formato_regional_comas(self):
        """Test que maneja correctamente números con comas (formato regional)."""
        cabecera = [("CODIGO PEDIDO", "P-12.48")]
        
        # Cantidad con coma como separador decimal
        detalle = [
            {"codigo": "R225.1012.BCO", "tiendas": {"ABC01": "10,0"}},
        ]
        
        excel_bytes = _crear_excel_v1(cabecera, detalle, tiendas=["ABC01"])
        
        service = PedidoService()
        resultado = service.parsear_excel_pedido(excel_bytes)
        
        # Debe convertir correctamente: "10,0" → 10
        assert len(resultado["data"]) == 1
        assert resultado["data"][0]["cantidad"] == 10
    
    def test_parsear_excel_cantidad_cero_ignora(self):
        """Test que ignora cantidades 0 (no las incluye en data)."""
        cabecera = [("CODIGO PEDIDO", "P-12.48")]
        
        detalle = [
            {"codigo": "R225.1012.BCO", "tiendas": {"ABC01": 0}},
        ]
        
        excel_bytes = _crear_excel_v1(cabecera, detalle, tiendas=["ABC01"])
        
        service = PedidoService()
        resultado = service.parsear_excel_pedido(excel_bytes)
        
        # Cantidades 0 no generan distribuciones
        distribuciones = [d for d in resultado["data"] if d["codigoProducto"] == "R225.1012.BCO"]
        assert len(distribuciones) == 0
    
    def test_parsear_excel_filas_vacias_ignora(self):
        """Test que ignora filas vacías y solo procesa filas con código de producto."""
        cabecera = [("CODIGO PEDIDO", "P-12.48")]
        
        detalle = [
            {"codigo": "R225.1012.BCO", "tiendas": {"ABC01": 10}},
            {"codigo": None, "tiendas": {}},       # Fila sin código (vacía)
            {"codigo": "", "tiendas": {}},          # Fila con código vacío
            {"codigo": "A889.3321.NEG", "tiendas": {"ABC02": 5}},
        ]
        
        excel_bytes = _crear_excel_v1(cabecera, detalle, tiendas=["ABC01", "ABC02"])
        
        service = PedidoService()
        resultado = service.parsear_excel_pedido(excel_bytes)
        
        # Solo debe procesar 2 filas válidas (las que tienen código de producto)
        assert resultado["resumen"]["filasValidas"] == 2
        assert len(resultado["data"]) == 2
    
    def test_parsear_excel_trim_espacios(self):
        """Test que hace trim de espacios en códigos de producto."""
        cabecera = [
            ("CODIGO PEDIDO", "  P-12.48  "),
        ]
        
        detalle = [
            {"codigo": "  R225.1012.BCO  ", "tiendas": {"ABC01": 10}},
        ]
        
        excel_bytes = _crear_excel_v1(cabecera, detalle, tiendas=["ABC01"])
        
        service = PedidoService()
        resultado = service.parsear_excel_pedido(excel_bytes)
        
        # Debe hacer trim del código de producto
        assert len(resultado["data"]) == 1
        assert resultado["data"][0]["codigoProducto"] == "R225.1012.BCO"
    
    def test_parsear_excel_case_insensitive_columnas(self):
        """Test que acepta columnas con diferentes mayúsculas/minúsculas."""
        cabecera = [("CODIGO PEDIDO", "P-12.48")]
        
        # Headers en minúsculas - el parser normaliza para buscar "codigo producto"
        headers = [
            "codigo producto", "", "", "", "", "", "", "", "",
            "", "", "",
            "ABC01"
        ]
        
        wb = Workbook()
        ws_cab = wb.active
        if ws_cab is None:
            raise RuntimeError("No se pudo crear hoja")
        ws_cab.title = "cabecera"
        ws_cab.cell(row=1, column=1, value="Campo")
        ws_cab.cell(row=1, column=2, value="Valor")
        for idx, (campo, valor) in enumerate(cabecera, start=2):
            ws_cab.cell(row=idx, column=1, value=campo)
            ws_cab.cell(row=idx, column=2, value=valor)
        
        ws_det = wb.create_sheet(title="detalle")
        # Fila 1: headers en minúsculas
        for col_idx, header in enumerate(headers, start=1):
            ws_det.cell(row=1, column=col_idx, value=header)
        # Fila 2: códigos
        ws_det.cell(row=2, column=13, value="ABC01")
        # Fila 3: datos
        ws_det.cell(row=3, column=1, value="R225.1012.BCO")
        ws_det.cell(row=3, column=13, value=10)
        
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        excel_bytes = output.read()
        
        service = PedidoService()
        resultado = service.parsear_excel_pedido(excel_bytes)
        
        # Debe parsear correctamente con headers en minúsculas
        assert len(resultado["data"]) == 1
        assert resultado["data"][0]["codigoProducto"] == "R225.1012.BCO"
    
    def test_parsear_excel_resumen_contadores(self):
        """Test que los contadores del resumen son correctos."""
        cabecera = [("CODIGO PEDIDO", "P-12.48")]
        
        detalle = [
            {"codigo": "R225.1012.BCO", "tiendas": {"ABC01": 10}},  # Válida
            {"codigo": "A889.3321.NEG", "tiendas": {"ABC02": 5}},   # Válida
            {"codigo": None, "tiendas": {}},                         # Inválida (sin código)
        ]
        
        excel_bytes = _crear_excel_v1(cabecera, detalle, tiendas=["ABC01", "ABC02"])
        
        service = PedidoService()
        resultado = service.parsear_excel_pedido(excel_bytes)
        
        # Verificar contadores (filasLeidas incluye todas las filas, filasValidas solo las con código)
        assert resultado["resumen"]["filasLeidas"] >= 2
        assert resultado["resumen"]["filasValidas"] == 2
    
    def test_parsear_excel_multiples_tiendas_por_producto(self):
        """Test que un producto con múltiples tiendas genera múltiples distribuciones."""
        cabecera = [("CODIGO PEDIDO", "P-12.48")]
        
        detalle = [
            {"codigo": "R225.1012.BCO", "tiendas": {"ABC01": 10, "ABC02": 5, "ABC03": 3}},
        ]
        
        excel_bytes = _crear_excel_v1(cabecera, detalle, tiendas=["ABC01", "ABC02", "ABC03"])
        
        service = PedidoService()
        resultado = service.parsear_excel_pedido(excel_bytes)
        
        # Un producto con 3 tiendas = 3 distribuciones
        assert len(resultado["data"]) == 3
        cantidades = sorted([d["cantidad"] for d in resultado["data"]])
        assert cantidades == [3, 5, 10]

    def test_parsear_excel_cantidad_decimal_preserva_decimales(self):
        """Test que cantidades decimales se preservan (soporte DECIMAL(18,4), no se truncan a entero)."""
        cabecera = [("CODIGO PEDIDO", "P-12.48")]

        detalle = [
            {"codigo": "R225.1012.BCO", "tiendas": {"ABC01": 10.7}},
        ]

        excel_bytes = _crear_excel_v1(cabecera, detalle, tiendas=["ABC01"])

        service = PedidoService()
        resultado = service.parsear_excel_pedido(excel_bytes)

        # Desde el soporte de cantidades decimales (columna DECIMAL(18,4)),
        # el parser ya no trunca a entero: preserva el valor tal cual.
        assert len(resultado["data"]) == 1
        assert resultado["data"][0]["cantidad"] == 10.7
